import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document(r'C:\Users\LENOVO\Downloads\Ke_hoach_thuc_hien_10_tuan_BloomGift.docx')
for i in range(len(doc.paragraphs)):
    p = doc.paragraphs[i]
    if p.text.strip():
        print(f"P{i}: text='{p.text}', align={p.alignment}, space_before={p.paragraph_format.space_before}, space_after={p.paragraph_format.space_after}, line_spacing={p.paragraph_format.line_spacing}")
        for r in p.runs:
            print(f"   run: '{r.text}', font={r.font.name}, size={r.font.size.pt if r.font.size else None}, bold={r.bold}, color={r.font.color.rgb if r.font.color else None}")

print("\n--- TABLE 1 CELL DETAILS ---")
t1 = doc.tables[0]
for r_idx, row in enumerate(t1.rows):
    for c_idx, cell in enumerate(row.cells):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = tcPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd')
        fill = shd.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill') if shd is not None else None
        p = cell.paragraphs[0]
        r = p.runs[0] if p.runs else None
        font_name = r.font.name if r else None
        font_size = r.font.size.pt if r and r.font.size else None
        bold = r.bold if r else None
        print(f"T1 R{r_idx}C{c_idx}: fill={fill}, text='{cell.text}', bold={bold}, size={font_size}")

print("\n--- TABLE 2 CELL DETAILS ---")
t2 = doc.tables[1]
for r_idx in [0, 1]:
    row = t2.rows[r_idx]
    for c_idx, cell in enumerate(row.cells):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = tcPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd')
        fill = shd.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill') if shd is not None else None
        p = cell.paragraphs[0]
        r = p.runs[0] if p.runs else None
        font_name = r.font.name if r else None
        font_size = r.font.size.pt if r and r.font.size else None
        bold = r.bold if r else None
        print(f"T2 R{r_idx}C{c_idx}: fill={fill}, text='{cell.text[:25]}', align={p.alignment}, bold={bold}, size={font_size}")
