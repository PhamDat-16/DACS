import os
import sys
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

if sys.stdout.encoding != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')

LOGO_PATH = os.path.join(os.getcwd(), "extracted_media", "word", "media", "image1.png")

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="000000", sz="8", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_header_table(doc, topic_text="Đề tài: Smart Parking System – Hệ thống quản lý bãi đỗ xe thông minh"):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="000000", sz="8", val="single")

    widths = [Cm(3.8), Cm(9.6), Cm(4.2)]

    # Cột 1: Logo
    cell0 = table.rows[0].cells[0]
    cell0.width = widths[0]
    set_cell_margins(cell0, top=120, bottom=120, left=120, right=120)
    p0 = cell0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(LOGO_PATH):
        r_logo = p0.add_run()
        r_logo.add_picture(LOGO_PATH, width=Inches(1.2))
    else:
        r_txt = p0.add_run("PHENIKAA\nUNIVERSITY")
        r_txt.font.name = 'Times New Roman'
        r_txt.font.size = Pt(10)
        r_txt.bold = True

    # Cột 2: Tiêu đề biên bản
    cell1 = table.rows[0].cells[1]
    cell1.width = widths[1]
    set_cell_margins(cell1, top=120, bottom=120, left=120, right=120)
    p1 = cell1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.line_spacing = 1.15
    p1.paragraph_format.space_before = Pt(2)
    p1.paragraph_format.space_after = Pt(2)

    r1_1 = p1.add_run("BIÊN BẢN HỌP NHÓM VÀ BÁO CÁO TIẾN ĐỘ CÔNG VIỆC\n")
    r1_1.font.name = 'Times New Roman'
    r1_1.font.size = Pt(11)
    r1_1.bold = True

    r1_2 = p1.add_run("(Học phần: Phát triển ứng dụng Hệ thống thông tin)\n")
    r1_2.font.name = 'Times New Roman'
    r1_2.font.size = Pt(9.5)
    r1_2.italic = True

    r1_3 = p1.add_run(topic_text)
    r1_3.font.name = 'Times New Roman'
    r1_3.font.size = Pt(10)
    r1_3.bold = True

    # Cột 3: Mã số BM
    cell2 = table.rows[0].cells[2]
    cell2.width = widths[2]
    set_cell_margins(cell2, top=120, bottom=120, left=120, right=120)
    p2 = cell2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p2.paragraph_format.line_spacing = 1.15
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(2)

    r2_1 = p2.add_run("Mã số: BM.CNTT.01.06\nLần ban hành: 01")
    r2_1.font.name = 'Times New Roman'
    r2_1.font.size = Pt(9.5)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(6)

def add_signature_table(doc, leader_name="Phạm Văn Đạt", instructor_name="ThS. Nguyễn Thị Thùy Liên"):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="FFFFFF", sz="0", val="none") # Không viền

    widths = [Cm(8.8), Cm(8.8)]

    # Cột 1: Giảng viên
    c0 = table.rows[0].cells[0]
    c0.width = widths[0]
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.line_spacing = 1.15
    r0 = p0.add_run(f"GIẢNG VIÊN HƯỚNG DẪN\n\n\n\n\n{instructor_name}\n(Ký và ghi rõ họ tên)")
    r0.font.name = 'Times New Roman'
    r0.font.size = Pt(11)
    r0.bold = True

    # Cột 2: Nhóm trưởng
    c1 = table.rows[0].cells[1]
    c1.width = widths[1]
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.line_spacing = 1.15
    r1 = p1.add_run(f"NHÓM TRƯỞNG\n\n\n\n\n{leader_name}\n(Ký và ghi rõ họ tên)")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11)
    r1.bold = True

def add_heading_section(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.bold = True
    return p

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.6)
    r_sym = p.add_run("• ")
    r_sym.font.name = 'Times New Roman'
    r_sym.font.size = Pt(10.5)
    r_sym.bold = True
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(10.5)
        r_pre.bold = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10.5)
    return p

def add_dash_bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.6)
    r_sym = p.add_run("– ")
    r_sym.font.name = 'Times New Roman'
    r_sym.font.size = Pt(10.5)
    r_sym.bold = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10.5)
    return p

# ==============================================================================
# TẠO FILE 1: BienBanHopNhom_SmartParking.docx (Giai đoạn Chuyển tiếp Tuần 4 -> Tuần 5)
# ==============================================================================
def create_doc_1():
    print("-> Đang tạo File 1: BienBanHopNhom_SmartParking.docx ...")
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.70)
        s.bottom_margin = Cm(1.70)
        s.left_margin = Cm(1.80)
        s.right_margin = Cm(1.80)
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)

    add_header_table(doc, "Đề tài: Smart Parking System – Hệ thống quản lý bãi đỗ xe thông minh")

    add_heading_section(doc, "1. Thời gian và địa điểm")
    add_bullet(doc, "20h00 – 21h30, ngày 20/09/2026.", bold_prefix="Thời gian: ")
    add_bullet(doc, "Họp trực tuyến qua ứng dụng Google Meet.", bold_prefix="Địa điểm/Hình thức: ")

    add_heading_section(doc, "2. Thành phần tham dự")
    add_bullet(doc, "Đầy đủ 03 thành viên thuộc Nhóm 4 nghiên cứu đề tài:")
    add_bullet(doc, "Phạm Văn Đạt – Trưởng nhóm / Chủ tọa cuộc họp (MSV: 23017150).")
    add_bullet(doc, "Nguyễn Duy Bảo – Kỹ sư Giao diện / Thư ký cuộc họp (MSV: [Bổ sung]).")
    add_bullet(doc, "Trần Tuấn Anh – Kỹ sư Thuật toán & AI / Thành viên (MSV: [Bổ sung]).")

    add_heading_section(doc, "3. Nội dung chi tiết")
    add_dash_bullet(doc, "Rà soát tiến độ Đồ án học phần Hệ thống thông tin với đề tài: Nghiên cứu, thiết kế và xây dựng Hệ thống quản lý bãi đỗ xe thông minh (Smart Parking System - SPS).")
    add_dash_bullet(doc, "Tổng hợp kết quả giai đoạn khảo sát hiện trạng và phân tích nghiệp vụ (BA): xác định rõ 3 tác nhân Customer / Staff / Admin, bảng quy tắc đặt chỗ, chống xung đột vé và yêu cầu phản hồi thời gian thực qua WebSocket.")
    add_dash_bullet(doc, "Rà soát toàn bộ tài liệu thiết kế hệ thống trước khi bắt tay vào lập trình: Sơ đồ Use Case tổng quát, Sơ đồ thực thể liên kết ERD chuẩn hóa 6 bảng (users, parking_lots, parking_slots, bookings, sessions, occupancy_history), thiết kế ma trận lưới bãi đỗ 2D, Activity Diagram và Sequence Diagram cho luồng Đặt chỗ và Quét QR Barrier.")
    add_dash_bullet(doc, "Thống nhất cấu trúc phân tầng công nghệ hiện đại: Backend FastAPI (Python 3.11, SQLAlchemy ORM, PostgreSQL) và Frontend Vue 3 (Composition API, Tailwind CSS, GSAP 3).")
    add_dash_bullet(doc, "Thống nhất quy tắc làm việc trên Git: mỗi thành viên thực hiện trên branch tính năng riêng (feature/auth, feature/booking, feature/astar), pull code mới nhất từ develop trước khi làm và bắt buộc tạo Pull Request kiểm tra chéo trước khi merge.")
    add_dash_bullet(doc, "Thống nhất kế hoạch tuần tiếp theo (Tuần 5): hoàn thiện khởi tạo dự án, kết nối CSDL PostgreSQL, hoàn thành xác thực JWT/Bcrypt và chuẩn bị cài đặt thuật toán A* và cơ chế khóa SELECT FOR UPDATE ở tuần 6.")

    add_heading_section(doc, "3.1. Thời gian thực hiện")
    add_bullet(doc, "Giai đoạn Tuần 4 – Tuần 5: từ 14/09/2026 đến 27/09/2026.")

    add_heading_section(doc, "3.2. Đánh giá tiến độ công việc")

    # Bảng đánh giá phân công thành viên
    headers_t2 = ["TT", "Mã sinh viên", "Họ và tên", "Nhiệm vụ phân công", "Nội dung đã hoàn thành", "Đánh giá tiến độ", "Ghi chú"]
    rows_t2 = [
        [
            "1",
            "23017150",
            "Phạm Văn Đạt",
            "Core Architecture – Database ERD – Authentication (JWT/Bcrypt)",
            "Đã hoàn thành thiết kế ERD 6 bảng, đặc tả kiến trúc 3 tầng và API Auth; bắt đầu khởi tạo khung dự án FastAPI và cấu hình PostgreSQL.",
            "Đúng tiến độ",
            "Tiếp tục hoàn thiện các endpoint đăng ký, đăng nhập và phân quyền role."
        ],
        [
            "2",
            "[Bổ sung]",
            "Nguyễn Duy Bảo",
            "Frontend UI/UX – Vue 3 Setup – Staff/Admin Dashboard",
            "Đã hoàn thiện bản vẽ Wireframe, thiết kế Sitemap điều hướng và cài đặt khung ứng dụng Vue 3 với Vite, Tailwind CSS và Pinia.",
            "Đúng tiến độ",
            "Tiếp tục dựng giao diện lưới bãi đỗ xe và component xe chạy GSAP 3."
        ],
        [
            "3",
            "[Bổ sung]",
            "Trần Tuấn Anh",
            "Pathfinding A* Algorithm – Booking Flow (SELECT FOR UPDATE)",
            "Đã nghiên cứu lý thuyết thuật toán A* trên lưới 2D và cơ chế khóa bản ghi tránh Race Condition; xây dựng ma trận lưới bãi đỗ tọa độ (x, y).",
            "Đúng tiến độ",
            "Chuẩn bị cài đặt thuật toán A* với heapq trong thư mục pathfinding/."
        ]
    ]

    t2 = doc.add_table(rows=len(rows_t2) + 1, cols=len(headers_t2))
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t2, color="000000", sz="8", val="single")

    col_w2 = [Cm(1.0), Cm(2.4), Cm(3.0), Cm(3.6), Cm(4.2), Cm(2.2), Cm(2.8)]

    for c_idx, heading in enumerate(headers_t2):
        cell = t2.rows[0].cells[c_idx]
        cell.width = col_w2[c_idx]
        set_cell_background(cell, "E7E6E6")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(heading)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.5)
        r.bold = True

    for r_idx, r_vals in enumerate(rows_t2):
        row = t2.rows[r_idx + 1]
        for c_idx, val in enumerate(r_vals):
            cell = row.cells[c_idx]
            cell.width = col_w2[c_idx]
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            if c_idx in [0, 1, 5]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9.0)

    p_sp_t2 = doc.add_paragraph()
    p_sp_t2.paragraph_format.space_after = Pt(4)

    add_heading_section(doc, "4. Kết luận và phân công sau cuộc họp")
    add_dash_bullet(doc, "Nhóm thống nhất tiến độ hiện tại hoàn toàn phù hợp và bám sát kế hoạch 10 tuần: giai đoạn phân tích nghiệp vụ và thiết kế kiến trúc hệ thống đã hoàn tất, chuyển trọng tâm sang hiện thực hóa mã nguồn.")
    add_dash_bullet(doc, "Các thành viên tiếp tục phát triển đúng module đã được phân công, chủ động trao đổi qua nhóm Zalo/Google Meet khi có điều chỉnh liên quan đến Schema cơ sở dữ liệu hoặc cấu trúc JSON Response của API.")
    add_dash_bullet(doc, "Mốc rà soát tiếp theo: hết ngày 27/09/2026; chuẩn bị nội dung họp báo cáo giai đoạn chuyển tiếp Tuần 6 – Tuần 7.")

    p_sp_sig = doc.add_paragraph()
    p_sp_sig.paragraph_format.space_after = Pt(8)

    add_signature_table(doc, leader_name="Phạm Văn Đạt")

    out_file1 = "BienBanHopNhom_SmartParking.docx"
    saved_file1 = out_file1
    try:
        doc.save(out_file1)
        print(f"==> Đã lưu File 1 tại: {os.path.abspath(out_file1)}")
    except PermissionError:
        saved_file1 = "BienBanHopNhom_SmartParking_Nhom4.docx"
        doc.save(saved_file1)
        print(f"==> File gốc đang mở trong WPS/Word. Đã lưu bản cập nhật tại: {os.path.abspath(saved_file1)}")

    user_dl = os.path.join(os.path.expanduser("~"), "Downloads")
    if os.path.exists(user_dl):
        dl_target = os.path.join(user_dl, out_file1)
        try:
            shutil.copy2(saved_file1, dl_target)
            print(f"==> Đã sao chép sang Downloads: {dl_target}")
        except PermissionError:
            dl_alt = os.path.join(user_dl, "BienBanHopNhom_SmartParking_Nhom4.docx")
            shutil.copy2(saved_file1, dl_alt)
            print(f"==> Bản trong Downloads đang mở. Đã lưu bản cập nhật vào: {dl_alt}")

# ==============================================================================
# TẠO FILE 2: Bien_ban_hop_nhom_SmartParking_Tuan_6_7_30-09-2026.docx
# ==============================================================================
def create_doc_2():
    print("-> Đang tạo File 2: Bien_ban_hop_nhom_SmartParking_Tuan_6_7_30-09-2026.docx ...")
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.70)
        s.bottom_margin = Cm(1.70)
        s.left_margin = Cm(1.80)
        s.right_margin = Cm(1.80)
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)

    add_header_table(doc, "Đề tài: Smart Parking System – Hệ thống quản lý bãi đỗ xe thông minh")

    add_heading_section(doc, "1. Thời gian và địa điểm")
    add_bullet(doc, "14h00 – 17h30, ngày 30/09/2026.", bold_prefix="Thời gian: ")
    add_bullet(doc, "Họp trực tuyến qua ứng dụng Google Meet.", bold_prefix="Địa điểm: ")

    add_heading_section(doc, "2. Thành phần tham dự")
    add_bullet(doc, "Đầy đủ 03 thành viên thuộc Nhóm 4 nghiên cứu đề tài:")
    add_bullet(doc, "Phạm Văn Đạt – Trưởng nhóm / Chủ tọa cuộc họp (MSV: 23017150).")
    add_bullet(doc, "Nguyễn Duy Bảo – Kỹ sư Giao diện / Thư ký cuộc họp (MSV: [Bổ sung]).")
    add_bullet(doc, "Trần Tuấn Anh – Kỹ sư Thuật toán & AI / Thành viên (MSV: [Bổ sung]).")

    add_heading_section(doc, "3. Nội dung chi tiết")
    add_bullet(doc, "Đánh giá tiến độ chuyển tiếp từ Tuần 5 sang Tuần 6 của đồ án Smart Parking System: Backend FastAPI đã hoàn thành kết nối PostgreSQL và xác thực JWT; giao diện tổng thể đã định hình.")
    add_bullet(doc, "Thống nhất nội dung triển khai trong Tuần 6: hoàn thiện thuật toán tìm đường tối ưu A* (app/pathfinding/astar.py) với cấu trúc dữ liệu heapq và hàm Heuristic Manhattan; kết nối bản đồ ParkingGrid.vue và hiệu ứng xe ảo chuyển động bằng GSAP 3.")
    add_bullet(doc, "Hoàn thiện cơ chế đặt giữ chỗ ô đỗ (Booking Flow): áp dụng kỹ thuật khóa dòng SELECT FOR UPDATE của SQLAlchemy để triệt tiêu tình trạng 2 khách đặt trùng ô đỗ (Double-booking); sinh mã vé điện tử QR Code động cho khách.")
    add_bullet(doc, "Chuẩn bị nội dung cho Tuần 7: phát triển Cổng nhân viên (Staff Portal) với phân quyền bãi làm việc (assigned_lot), kết nối WebSocket cập nhật trạng thái ô đỗ thời gian thực, tích hợp quét mã QR bằng webcam/camera và cửa sổ điều khiển cần Barrier mở tự động.")
    add_bullet(doc, "Thống nhất quy trình làm việc trên Git: mỗi thành viên làm việc trên branch riêng, cập nhật từ develop trước khi làm và tạo Pull Request khi hoàn thành phần việc.")
    add_bullet(doc, "Rà soát sự liên kết dữ liệu giữa Booking – ParkingSlot – Session để tránh sai lệch dữ liệu khi tích hợp luồng Check-in và Check-out.")

    add_heading_section(doc, "3.1. Thời gian thực hiện")
    add_bullet(doc, "Giai đoạn Tuần 6 – Tuần 7: từ 28/09/2026 đến 11/10/2026.")

    add_heading_section(doc, "3.2. Đánh giá tiến độ công việc")

    # Bảng đánh giá chi tiết tuần 6-7
    headers_t2 = ["TT", "Mã sinh viên", "Họ và tên", "Nhiệm vụ phân công", "Nội dung đã/đang thực hiện", "Đánh giá tiến độ", "Ghi chú"]
    rows_t2 = [
        [
            "1",
            "23017150",
            "Phạm Văn Đạt",
            "Booking Flow – SELECT FOR UPDATE – QR Code – WebSocket Manager",
            "Tiếp tục hoàn thiện API Đặt chỗ có khóa SELECT FOR UPDATE; kiểm tra sinh mã QR Code cho vé xe; viết ConnectionManager phát sóng WebSocket khi ô đỗ đổi trạng thái.",
            "Đúng kế hoạch",
            "Phối hợp tích hợp bản đồ ô đỗ với Frontend và chuẩn bị test đa luồng."
        ],
        [
            "2",
            "[Bổ sung]",
            "Nguyễn Duy Bảo",
            "GSAP Animation – Customer Parking Grid – Staff Portal UI",
            "Xây dựng component bản đồ ParkingGrid.vue và hiệu ứng xe ảo GSAP 3; thiết kế giao diện Staff Portal và tích hợp thư viện quét mã QR qua webcam.",
            "Đúng kế hoạch",
            "Kiểm tra hiển thị đường nét đứt phát sáng và xoay đầu xe theo góc cua."
        ],
        [
            "3",
            "[Bổ sung]",
            "Trần Tuấn Anh",
            "A* Algorithm (heapq) – Pipeline dữ liệu – Random Forest Model",
            "Hoàn thiện hàm tìm đường A* né vật cản tĩnh/động; xây dựng pipeline sinh 4.320 mẫu dữ liệu lịch sử và huấn luyện mô hình Random Forest 200 cây quyết định.",
            "Đúng kế hoạch",
            "Kiểm thử đánh giá độ chính xác MAE ~ 0.058 và tích hợp file mô hình .joblib."
        ]
    ]

    t2 = doc.add_table(rows=len(rows_t2) + 1, cols=len(headers_t2))
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t2, color="000000", sz="8", val="single")

    col_w2 = [Cm(1.0), Cm(2.2), Cm(3.0), Cm(3.8), Cm(4.6), Cm(2.2), Cm(2.4)]

    for c_idx, heading in enumerate(headers_t2):
        cell = t2.rows[0].cells[c_idx]
        cell.width = col_w2[c_idx]
        set_cell_background(cell, "E7E6E6")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(heading)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.5)
        r.bold = True

    for r_idx, r_vals in enumerate(rows_t2):
        row = t2.rows[r_idx + 1]
        for c_idx, val in enumerate(r_vals):
            cell = row.cells[c_idx]
            cell.width = col_w2[c_idx]
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            if c_idx in [0, 1, 5]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9.0)

    p_sp_t2 = doc.add_paragraph()
    p_sp_t2.paragraph_format.space_after = Pt(4)

    add_heading_section(doc, "4. Kế hoạch công việc sau cuộc họp")
    add_bullet(doc, "Hoàn thành toàn bộ các nội dung Tuần 6 (A* Navigation, Đặt chỗ khóa hàng, Vé QR) trước khi chuyển sang triển khai đầy đủ cổng Barrier và WebSocket ở Tuần 7.")
    add_bullet(doc, "Tích hợp dữ liệu giữa Booking, ParkingSlot và Session; kiểm tra tính nhất quán thông tin biển số xe, thời gian bắt đầu và kết thúc.")
    add_bullet(doc, "Trong Tuần 7, ưu tiên hoàn thiện Cổng nhân viên (Staff Portal), quét mã QR mở Barrier, cập nhật trạng thái ô đỗ thời gian thực và tự động tính tiền.")
    add_bullet(doc, "Mỗi thành viên cập nhật tiến độ trên branch cá nhân và chỉ merge vào develop sau khi đã tự kiểm thử cục bộ.")
    add_bullet(doc, "Chuẩn bị minh chứng tiến độ (ảnh chụp màn hình kết quả chạy thực tế) để phục vụ buổi báo cáo tiến độ học phần tiếp theo.")

    add_heading_section(doc, "5. Kết luận cuộc họp")
    p_conclusion = doc.add_paragraph()
    p_conclusion.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_conclusion.paragraph_format.space_before = Pt(0)
    p_conclusion.paragraph_format.space_after = Pt(8)
    p_conclusion.paragraph_format.line_spacing = 1.15
    r_c = p_conclusion.add_run("Nhóm thống nhất tiếp tục triển khai theo đúng kế hoạch 10 tuần đã đề ra. Giai đoạn Tuần 6 tập trung hoàn thiện thuật toán tìm đường A* và luồng Đặt giữ chỗ an toàn không xung đột; Tuần 7 chuyển trọng tâm sang Cổng nhân viên, vận hành Barrier qua quét mã QR, truyền thông WebSocket thời gian thực và quản lý doanh thu. Các thành viên phối hợp chặt chẽ theo từng module để bảo đảm dữ liệu được tích hợp đồng bộ và tránh xung đột mã nguồn.")
    r_c.font.name = 'Times New Roman'
    r_c.font.size = Pt(10.5)

    p_sp_sig = doc.add_paragraph()
    p_sp_sig.paragraph_format.space_after = Pt(8)

    add_signature_table(doc, leader_name="Phạm Văn Đạt")

    out_file2 = "Bien_ban_hop_nhom_SmartParking_Tuan_6_7_30-09-2026.docx"
    saved_file2 = out_file2
    try:
        doc.save(out_file2)
        print(f"==> Đã lưu File 2 tại: {os.path.abspath(out_file2)}")
    except PermissionError:
        saved_file2 = "Bien_ban_hop_nhom_SmartParking_Tuan_6_7_30-09-2026_Nhom4.docx"
        doc.save(saved_file2)
        print(f"==> File gốc đang mở trong WPS/Word. Đã lưu bản cập nhật tại: {os.path.abspath(saved_file2)}")

    user_dl = os.path.join(os.path.expanduser("~"), "Downloads")
    if os.path.exists(user_dl):
        dl_target = os.path.join(user_dl, out_file2)
        try:
            shutil.copy2(saved_file2, dl_target)
            print(f"==> Đã sao chép sang Downloads: {dl_target}")
        except PermissionError:
            dl_alt = os.path.join(user_dl, "Bien_ban_hop_nhom_SmartParking_Tuan_6_7_30-09-2026_Nhom4.docx")
            shutil.copy2(saved_file2, dl_alt)
            print(f"==> Bản trong Downloads đang mở. Đã lưu bản cập nhật vào: {dl_alt}")

if __name__ == "__main__":
    create_doc_1()
    create_doc_2()
    print("=== HOÀN TẤT TẠO 2 BIÊN BẢN HỌP NHÓM! ===")

