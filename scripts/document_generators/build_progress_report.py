import subprocess
import os
import sys
import base64
import shutil
import pypdf
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

PROJECT_DIR = os.path.abspath(r"D:\smart-parking-system")
HTML_FILE = os.path.join(PROJECT_DIR, "Bao_cao_tien_do_du_an_Smart_Parking_System.html")
PDF_FILE = os.path.join(PROJECT_DIR, "Bao_cao_tien_do_du_an_Smart_Parking_System.pdf")
DOCX_FILE = os.path.join(PROJECT_DIR, "Bao_cao_tien_do_du_an_Smart_Parking_System.docx")

USER_DOWNLOADS = os.path.join(os.path.expanduser("~"), "Downloads")
DOWNLOAD_PDF = os.path.join(USER_DOWNLOADS, "Bao_cao_tien_do_du_an_Smart_Parking_System.pdf")
DOWNLOAD_DOCX = os.path.join(USER_DOWNLOADS, "Bao_cao_tien_do_du_an_Smart_Parking_System.docx")

EDGE_PATH = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
IMG_DIR = r"C:\Users\LENOVO\.gemini\antigravity\brain\34425343-ff1b-46ef-a476-294ad3a89c79"

def get_base64_img(filename):
    path = os.path.join(IMG_DIR, filename)
    if not os.path.exists(path):
        path = os.path.join(PROJECT_DIR, filename)
    if os.path.exists(path):
        with open(path, "rb") as f:
            data = base64.b64encode(f.read()).decode("utf-8")
            return f"data:image/png;base64,{data}"
    return ""

img_staff_checkin = get_base64_img("staff_checkin_map_synced.png")
img_staff_checkout = get_base64_img("staff_checkout_selected_vehicle.png")
img_customer_dash = get_base64_img("customer_dashboard_active_session.png")
img_customer_receipt = get_base64_img("customer_history_receipt_modal.png")
img_admin_passes = get_base64_img("admin_monthly_passes_table.png")
img_dup_warning = get_base64_img("duplicate_checkin_warning.png")

# ==============================================================================
# 1. TẠO TỆP HTML ĐỒNG THỜI XUẤT RA PDF QUA MICROSOFT EDGE HEADLESS
# ==============================================================================
html_content = f"""<!DOCTYPE html>
<html lang="vi">
<head>
    <meta charset="UTF-8">
    <title>Báo Cáo Tiến Độ Dự Án - Smart Parking System</title>
    <style>
        @page {{
            size: A4;
            margin: 10mm 14mm 10mm 14mm;
        }}
        * {{
            box-sizing: border-box;
            -webkit-print-color-adjust: exact !important;
            print-color-adjust: exact !important;
        }}
        body {{
            font-family: 'Segoe UI', -apple-system, BlinkMacSystemFont, 'Times New Roman', Roboto, Arial, sans-serif;
            color: #1e293b;
            line-height: 1.33;
            font-size: 9.2pt;
            margin: 0;
            padding: 0;
            background-color: #ffffff;
        }}
        
        /* HEADER BÁO CÁO HỌC THUẬT CHUẨN MỰC */
        .academic-header {{
            border-bottom: 2px solid #1e3a8a;
            padding-bottom: 5px;
            margin-bottom: 8px;
            display: flex;
            justify-content: space-between;
            align-items: flex-end;
        }}
        .univ-title {{
            font-size: 8.5pt;
            font-weight: 700;
            color: #475569;
            text-transform: uppercase;
            letter-spacing: 0.3px;
            margin-bottom: 2px;
        }}
        .doc-title {{
            font-size: 15pt;
            font-weight: 900;
            color: #1e3a8a;
            margin: 0;
            letter-spacing: -0.2px;
            text-transform: uppercase;
        }}
        .doc-subtitle {{
            font-size: 9.8pt;
            color: #0f172a;
            margin-top: 2px;
            font-weight: 700;
        }}
        .meta-box {{
            text-align: right;
            font-size: 8.5pt;
            color: #334155;
            line-height: 1.35;
        }}
        
        .badge {{
            display: inline-block;
            padding: 2px 7px;
            border-radius: 9999px;
            font-size: 8pt;
            font-weight: 800;
            text-transform: uppercase;
            letter-spacing: 0.3px;
        }}
        .badge-success {{
            background-color: #dcfce7;
            color: #15803d;
            border: 1px solid #86efac;
        }}

        /* KPI SUMMARY CARDS */
        .kpi-row {{
            display: grid;
            grid-template-columns: repeat(5, 1fr);
            gap: 7px;
            margin-bottom: 8px;
        }}
        .kpi-box {{
            border: 1px solid #cbd5e1;
            border-top: 3px solid #1e3a8a;
            background: #f8fafc;
            border-radius: 5px;
            padding: 5px 4px;
            text-align: center;
        }}
        .kpi-box.green {{ border-top-color: #10b981; }}
        .kpi-box.purple {{ border-top-color: #8b5cf6; }}
        .kpi-box.amber {{ border-top-color: #f59e0b; }}
        .kpi-box.blue {{ border-top-color: #0284c7; }}

        .kpi-num {{
            font-size: 14pt;
            font-weight: 900;
            color: #0f172a;
            line-height: 1.1;
        }}
        .kpi-label {{
            font-size: 7pt;
            font-weight: 700;
            color: #64748b;
            text-transform: uppercase;
            margin-top: 2px;
        }}

        .section-title {{
            font-size: 10pt;
            font-weight: 800;
            color: #0f172a;
            text-transform: uppercase;
            letter-spacing: 0.2px;
            margin: 7px 0 5px 0;
            padding-left: 7px;
            border-left: 3px solid #1e3a8a;
            display: flex;
            align-items: center;
            justify-content: space-between;
        }}
        .section-title span.tag {{
            font-size: 7.5pt;
            font-weight: 700;
            color: #15803d;
            background: #dcfce7;
            padding: 1px 5px;
            border-radius: 4px;
            border: 1px solid #bbf7d0;
            text-transform: none;
        }}

        /* BẢNG BIỂU CHUẨN FORMAL */
        table {{
            width: 100%;
            border-collapse: collapse;
            margin-bottom: 7px;
            font-size: 8.5pt;
        }}
        th, td {{
            padding: 4px 6px;
            text-align: left;
            border: 1px solid #cbd5e1;
        }}
        th {{
            background-color: #f1f5f9;
            color: #1e293b;
            font-weight: 800;
            font-size: 8pt;
            text-transform: uppercase;
            letter-spacing: 0.2px;
        }}
        tr:nth-child(even) {{
            background-color: #f8fafc;
        }}

        .progress-bar-bg {{
            background-color: #e2e8f0;
            border-radius: 3px;
            height: 5px;
            width: 50px;
            overflow: hidden;
            display: inline-block;
            vertical-align: middle;
            margin-right: 3px;
        }}
        .progress-bar-fill {{
            background-color: #10b981;
            height: 100%;
            border-radius: 3px;
        }}

        .module-card {{
            background: #ffffff;
            border: 1px solid #cbd5e1;
            border-radius: 5px;
            padding: 5px 8px;
            margin-bottom: 5px;
        }}
        .module-header {{
            display: flex;
            justify-content: space-between;
            align-items: center;
            margin-bottom: 2px;
        }}
        .module-name {{
            font-weight: 800;
            font-size: 9pt;
            color: #1e3a8a;
        }}
        .module-desc {{
            color: #334155;
            font-size: 8.5pt;
            margin: 0;
            text-align: justify;
            line-height: 1.3;
        }}

        /* HÌNH ẢNH MINH CHỨNG THỰC NGHIỆM */
        .img-grid-2 {{
            display: grid;
            grid-template-columns: 1fr 1fr;
            gap: 7px;
            margin-bottom: 6px;
        }}
        .img-card {{
            border: 1px solid #cbd5e1;
            border-radius: 5px;
            overflow: hidden;
            background: #0f172a;
            box-shadow: 0 1px 2px rgba(0,0,0,0.05);
        }}
        .img-card img {{
            width: 100%;
            height: 162px;
            object-fit: cover;
            display: block;
        }}
        .img-caption {{
            font-size: 7.8pt;
            color: #334155;
            padding: 3px 5px;
            background: #f8fafc;
            border-top: 1px solid #e2e8f0;
            text-align: center;
            font-weight: 600;
        }}

        .page-break {{
            page-break-before: always;
            break-before: page;
        }}

        .footer-sign {{
            margin-top: 25px;
            display: grid;
            grid-template-columns: 1fr 1fr;
            text-align: center;
            font-size: 9.2pt;
        }}
        .sign-title {{
            font-weight: 800;
            margin-bottom: 40px;
        }}
        .sign-name {{
            font-weight: 800;
            color: #0f172a;
        }}
    </style>
</head>
<body>

    <!-- ==================== TRANG 1 ==================== -->
    <div class="academic-header">
        <div>
            <div class="univ-title">TRƯỜNG ĐẠI HỌC PHENIKAA – KHOA KHOA HỌC MÁY TÍNH</div>
            <div class="doc-title">BÁO CÁO TIẾN ĐỘ DỰ ÁN</div>
            <div class="doc-subtitle">HỆ THỐNG QUẢN LÝ VÀ DẪN ĐƯỜNG BÃI ĐỖ XE THÔNG MINH (SMART PARKING SYSTEM - SPS)</div>
        </div>
        <div class="meta-box">
            <div><strong>Học phần:</strong> CSE702131-1-1-26(N01)</div>
            <div><strong>GVHD:</strong> ThS. Nguyễn Thị Thùy Liên</div>
            <div><strong>Nhóm sinh viên:</strong> Nhóm 4 (03 thành viên)</div>
            <div><strong>Đánh giá:</strong> <span class="badge badge-success">HOÀN THÀNH 100% MỤC TIÊU</span></div>
        </div>
    </div>

    <!-- 5 KPI CARDS -->
    <div class="kpi-row">
        <div class="kpi-box green">
            <div class="kpi-num">100%</div>
            <div class="kpi-label">Tiến Độ Toàn Dự Án</div>
        </div>
        <div class="kpi-box blue">
            <div class="kpi-num">28 / 28</div>
            <div class="kpi-label">API Endpoints Hoàn Tất</div>
        </div>
        <div class="kpi-box purple">
            <div class="kpi-num">&lt; 5 ms</div>
            <div class="kpi-label">Độ Trễ Dẫn Đường A*</div>
        </div>
        <div class="kpi-box amber">
            <div class="kpi-num">0.058</div>
            <div class="kpi-label">Sai Số AI (MAE &lt; 0.06)</div>
        </div>
        <div class="kpi-box green">
            <div class="kpi-num">100%</div>
            <div class="kpi-label">Unit & Concurrency Tests</div>
        </div>
    </div>

    <div class="section-title">
        <span>I. BẢNG TIẾN ĐỘ CHI TIẾT TỪNG PHÂN HỆ NGHIỆP VỤ (WBS - NHÓM 4)</span>
        <span class="tag">Giai đoạn nghiệm thu &amp; bảo vệ</span>
    </div>

    <table>
        <thead>
            <tr>
                <th style="width: 25%;">Phân hệ chức năng</th>
                <th style="width: 47%;">Các hạng mục kỹ thuật đã hoàn thiện &amp; tích hợp</th>
                <th style="width: 14%; text-align: center;">Thành viên</th>
                <th style="width: 14%; text-align: center;">Tiến độ</th>
            </tr>
        </thead>
        <tbody>
            <tr>
                <td><strong>1. Backend Core &amp; CSDL</strong></td>
                <td>6 bảng CSDL quan hệ chuẩn 3NF trên PostgreSQL, Alembic Migration, Bcrypt, JWT Bearer Token, WebSocket realtime, Khóa dòng SELECT FOR UPDATE.</td>
                <td style="text-align: center; font-size: 8pt;">Phạm Văn Đạt</td>
                <td style="text-align: center;">
                    <span class="progress-bar-bg"><span class="progress-bar-fill" style="width: 100%;"></span></span> 100%
                </td>
            </tr>
            <tr>
                <td><strong>2. Thuật toán Dẫn đường A*</strong></td>
                <td>Ma trận lưới tọa độ 2D, hàm Heuristic Manhattan, Priority Queue né vật cản tĩnh/động, API GET /api/lots/{{id}}/navigate phản hồi &lt; 5ms.</td>
                <td style="text-align: center; font-size: 8pt;">Trần Tuấn Anh</td>
                <td style="text-align: center;">
                    <span class="progress-bar-bg"><span class="progress-bar-fill" style="width: 100%;"></span></span> 100%
                </td>
            </tr>
            <tr>
                <td><strong>3. AI Engine Dự báo lấp đầy</strong></td>
                <td>Sinh 4.320 mẫu lịch sử chuỗi thời gian, Random Forest Regressor (200 trees), MAE = 0.058, nạp Singleton model, API dự báo 24h kèm scrubber.</td>
                <td style="text-align: center; font-size: 8pt;">Trần Tuấn Anh</td>
                <td style="text-align: center;">
                    <span class="progress-bar-bg"><span class="progress-bar-fill" style="width: 100%;"></span></span> 100%
                </td>
            </tr>
            <tr>
                <td><strong>4. Cổng Khách hàng (Customer)</strong></td>
                <td>Giao diện Vue 3 Dark Cyber, sơ đồ lưới 2D ParkingGrid, diễn hoạt xe chạy GSAP 3, đặt chỗ an toàn, vé điện tử QR Code động, thẻ tháng cư dân.</td>
                <td style="text-align: center; font-size: 8pt;">Nguyễn Duy Bảo</td>
                <td style="text-align: center;">
                    <span class="progress-bar-bg"><span class="progress-bar-fill" style="width: 100%;"></span></span> 100%
                </td>
            </tr>
            <tr>
                <td><strong>5. Cổng Nhân viên (Staff Portal)</strong></td>
                <td>Trực barrier theo bãi phân công (assigned_lot), check-in quét QR &amp; vãng lai, cảnh báo check-in trùng lặp, check-out tính tiền theo giờ &amp; mở barrier.</td>
                <td style="text-align: center; font-size: 8pt;">Đạt &amp; Bảo</td>
                <td style="text-align: center;">
                    <span class="progress-bar-bg"><span class="progress-bar-fill" style="width: 100%;"></span></span> 100%
                </td>
            </tr>
            <tr>
                <td><strong>6. Cổng Quản trị viên (Admin)</strong></td>
                <td>Sơ đồ slot Inspector (Bảo trì/Khả dụng), tạo tài khoản nhân viên phân theo bãi xe, duyệt vé tháng cư dân, báo cáo phân tích doanh thu &amp; ca trực.</td>
                <td style="text-align: center; font-size: 8pt;">Cả nhóm</td>
                <td style="text-align: center;">
                    <span class="progress-bar-bg"><span class="progress-bar-fill" style="width: 100%;"></span></span> 100%
                </td>
            </tr>
        </tbody>
    </table>

    <div class="section-title">
        <span>II. CHI TIẾT CÁC ĐẶC TẢ KỸ THUẬT VỪA HOÀN THÀNH VÀ NÂNG CẤP</span>
    </div>

    <div class="module-card">
        <div class="module-header">
            <span class="module-name">1. Phân quyền nhân viên theo cơ sở công tác (Assigned Lot) &amp; Đăng nhập linh hoạt</span>
            <span class="badge badge-success">Đã hoàn thiện</span>
        </div>
        <p class="module-desc">
            Bổ sung trường dữ liệu <code>assigned_lot</code> trên CSDL và giao diện Admin, cho phép Quản trị viên tạo tài khoản nhân viên gắn liền với từng bãi xe cụ thể (Bãi Đỗ Xe KĐT Đô Nghĩa, Vincom Center, Aeon Mall...). Cải tiến cơ chế xác thực: hỗ trợ đăng nhập linh hoạt bằng cả <strong>Tên đăng nhập (Username: admin, staff_baove...) HOẶC Email</strong>, tự động chuẩn hóa chuỗi và nhận diện bí danh thông minh.
        </p>
    </div>

    <div class="module-card">
        <div class="module-header">
            <span class="module-name">2. Quy trình Check-in/Check-out toàn diện &amp; Chống trùng lặp (Anti Duplicate Check-in)</span>
            <span class="badge badge-success">Đã hoàn thiện</span>
        </div>
        <p class="module-desc">
            Hoàn thiện quy trình xuất nhập bến xe cho Nhân viên trực: Tự động phát hiện và cảnh báo chặn khi một xe cố tình check-in lần hai khi chưa xuất bến. Khi Check-out: hệ thống tự động tính cước gửi xe theo thời gian thực tế, chuyển trạng thái phiên thành <code>completed</code> và <strong>giải phóng slot đỗ xe từ Đang đỗ (Đỏ) về Trống (Xanh) tức thời</strong> trên cả CSDL và sơ đồ WebSocket.
        </p>
    </div>

    <div class="module-card">
        <div class="module-header">
            <span class="module-name">3. Quản lý Vé tháng cư dân (Monthly Passes) &amp; Hóa đơn điện tử</span>
            <span class="badge badge-success">Đã hoàn thiện</span>
        </div>
        <p class="module-desc">
            Tích hợp đầy đủ luồng vé tháng: Khách hàng đăng ký gói vé tháng theo biển số, Admin tiếp nhận và kích hoạt vé; khách hàng thanh toán trực tuyến và lưu trữ thẻ vé điện tử. Mỗi lượt gửi xe hoàn tất đều phát sinh hóa đơn chi tiết (Receipt Modal) hiển thị biển số, giờ vào, giờ ra, tổng tiền và mã giao dịch.
        </p>
    </div>

    <!-- ==================== TRANG 2 ==================== -->
    <div class="page-break"></div>

    <div class="academic-header">
        <div>
            <div class="univ-title">TRƯỜNG ĐẠI HỌC PHENIKAA – KHOA KHOA HỌC MÁY TÍNH</div>
            <div class="doc-title">HÌNH ẢNH MINH CHỨNG THỰC NGHIỆM TỪ HỆ THỐNG</div>
            <div class="doc-subtitle">DỰ ÁN SMART PARKING SYSTEM (SPS) – NHÓM 4 – GVHD: THS. NGUYỄN THỊ THÙY LIÊN</div>
        </div>
        <div class="meta-box">
            <div><strong>Trang:</strong> 2 / 3</div>
            <div><strong>Môi trường:</strong> Full-Stack Live System</div>
        </div>
    </div>

    <div class="section-title">
        <span>III. HÌNH ẢNH MINH CHỨNG THỰC TẾ TRÊN CÁC PHÂN HỆ VẬN HÀNH</span>
    </div>

    <div class="img-grid-2">
        <div class="img-card">
            <img src="{img_staff_checkin}" alt="Check-in Map">
            <div class="img-caption">Hình 1: Nhân viên trực bãi Đô Nghĩa – Sơ đồ ô đỗ thời gian thực &amp; Cổng nhận xe</div>
        </div>
        <div class="img-card">
            <img src="{img_staff_checkout}" alt="Check-out Selected Vehicle">
            <div class="img-caption">Hình 2: Nhân viên Check-out – Tính cước theo thời gian đỗ &amp; Mở barrier xuất bến</div>
        </div>
        <div class="img-card">
            <img src="{img_customer_dash}" alt="Customer Active Session">
            <div class="img-caption">Hình 3: Khách hàng Dashboard – Theo dõi thời gian phiên đỗ xe &amp; Bản đồ dẫn đường</div>
        </div>
        <div class="img-card">
            <img src="{img_customer_receipt}" alt="Customer Receipt Modal">
            <div class="img-caption">Hình 4: Hóa đơn điện tử xuất bến – Chi tiết thời gian gửi, cước phí &amp; Thanh toán</div>
        </div>
        <div class="img-card">
            <img src="{img_admin_passes}" alt="Admin Monthly Passes">
            <div class="img-caption">Hình 5: Quản trị viên – Quản lý danh sách và phê duyệt vé tháng cư dân</div>
        </div>
        <div class="img-card">
            <img src="{img_dup_warning}" alt="Duplicate Checkin Warning">
            <div class="img-caption">Hình 6: Cơ chế an toàn – Cảnh báo chặn xe đã trong bãi cố tình check-in trùng lặp</div>
        </div>
    </div>

    <!-- ==================== TRANG 3 ==================== -->
    <div class="page-break"></div>

    <div class="academic-header">
        <div>
            <div class="univ-title">TRƯỜNG ĐẠI HỌC PHENIKAA – KHOA KHOA HỌC MÁY TÍNH</div>
            <div class="doc-title">BẢO ĐẢM CHẤT LƯỢNG VÀ KẾ HOẠCH BẢO VỆ ĐỒ ÁN</div>
            <div class="doc-subtitle">DỰ ÁN SMART PARKING SYSTEM (SPS) – NHÓM 4 – THÁNG 10/2026</div>
        </div>
        <div class="meta-box">
            <div><strong>Trang:</strong> 3 / 3</div>
            <div><strong>Giai đoạn:</strong> Đã sẵn sàng nghiệm thu</div>
        </div>
    </div>

    <div class="section-title">
        <span>IV. KẾT QUẢ KIỂM THỬ TỰ ĐỘNG &amp; BẢO ĐẢM CHẤT LƯỢNG (QA)</span>
    </div>

    <table>
        <thead>
            <tr>
                <th style="width: 30%;">Bộ kiểm thử (Test Suite)</th>
                <th style="width: 44%;">Mục tiêu kiểm định</th>
                <th style="width: 13%; text-align: center;">Số lượng</th>
                <th style="width: 13%; text-align: center;">Kết quả</th>
            </tr>
        </thead>
        <tbody>
            <tr>
                <td><strong>Unit Test: Thuật toán A*</strong> (<code>test_astar.py</code>)</td>
                <td>Đường đi thẳng, tự động né vật cản tĩnh/động, điểm xuất phát trùng đích, ngõ cụt.</td>
                <td style="text-align: center;">6 / 6 cases</td>
                <td style="text-align: center;"><span class="badge badge-success">100% PASSED</span></td>
            </tr>
            <tr>
                <td><strong>Unit Test: AI Engine</strong> (<code>test_ml_engine.py</code>)</td>
                <td>Dữ liệu phân phối chuẩn, tính tái lập theo seed, hội tụ mô hình, miền dự báo [0, 1].</td>
                <td style="text-align: center;">5 / 5 cases</td>
                <td style="text-align: center;"><span class="badge badge-success">100% PASSED</span></td>
            </tr>
            <tr>
                <td><strong>Concurrency Test: Race Condition</strong> (<code>test_double_booking.py</code>)</td>
                <td>Mô phỏng 50 tiến trình đồng thời đặt cùng 1 slot tại cùng 1 ms (SELECT FOR UPDATE).</td>
                <td style="text-align: center;">50 workers</td>
                <td style="text-align: center;"><span class="badge badge-success">TRIỆT TIÊU XUNG ĐỘT</span></td>
            </tr>
            <tr>
                <td><strong>E2E Playwright Automation</strong> (<code>scratch/*.py</code>)</td>
                <td>Luồng liên hoàn: Đăng nhập -> Check-in -> Giải phóng slot -> Xuất bến tính cước.</td>
                <td style="text-align: center;">6 kịch bản</td>
                <td style="text-align: center;"><span class="badge badge-success">ỔN ĐỊNH 100%</span></td>
            </tr>
        </tbody>
    </table>

    <div class="section-title">
        <span>V. ĐÁNH GIÁ RỦI RO KỸ THUẬT &amp; GIẢI PHÁP ĐÃ ÁP DỤNG</span>
    </div>

    <table>
        <thead>
            <tr>
                <th style="width: 25%;">Hạng mục rủi ro</th>
                <th style="width: 35%;">Khả năng tác động &amp; Nguy cơ</th>
                <th style="width: 40%;">Giải pháp công nghệ đã triển khai</th>
            </tr>
        </thead>
        <tbody>
            <tr>
                <td><strong>Xung đột đặt chỗ (Race Condition)</strong></td>
                <td>Nhiều khách cùng bấm đặt 1 vị trí đỗ tại cùng 1 thời điểm dẫn đến trùng lặp dữ liệu.</td>
                <td>Cài đặt khóa cấp dòng <code>SELECT FOR UPDATE</code> trên PostgreSQL, trả về HTTP 400/409 an toàn cho tiến trình sau.</td>
            </tr>
            <tr>
                <td><strong>Độ trễ tính toán tìm đường A*</strong></td>
                <td>Ma trận bãi xe kích thước lớn khiến thuật toán phản hồi chậm ảnh hưởng trải nghiệm lái xe.</td>
                <td>Tối ưu hàm ước lượng Heuristic Manhattan và hàng đợi ưu tiên <code>heapq</code>, rút ngắn thời gian xử lý xuống &lt; 5ms.</td>
            </tr>
            <tr>
                <td><strong>Xe gian lận Check-in trùng lặp</strong></td>
                <td>Tài xế cố tình quét mã lần 2 khi chưa xuất bến gây lỗi đếm thừa vị trí đỗ.</td>
                <td>Khóa logic: Kiểm tra phiên đỗ đang active; nếu tồn tại xe chưa checkout sẽ từ chối và cảnh báo âm thanh/thông báo.</td>
            </tr>
        </tbody>
    </table>

    <div class="section-title">
        <span>VI. HỒ SƠ DỰ ÁN VÀ KẾ HOẠCH BẢO VỆ ĐỒ ÁN</span>
    </div>

    <ul style="margin: 4px 0 12px 0; padding-left: 18px; font-size: 8.8pt; color: #334155; line-height: 1.4;">
        <li><strong>Báo cáo thuyết minh đồ án:</strong> Đã hoàn thành toàn diện 85+ trang chuẩn học thuật Phenikaa (<code>BAO_CAO_DACS_Smart_Parking_System.docx</code>).</li>
        <li><strong>Hồ sơ nghiệm thu:</strong> Đầy đủ Kế hoạch 10 tuần, 2 Biên bản họp nhóm, Bộ 50 câu hỏi &amp; trả lời bảo vệ đồ án, Kịch bản thuyết trình và hướng dẫn vận hành 1-click (<code>run_web.bat</code>).</li>
        <li><strong>Trạng thái bảo vệ:</strong> Hệ thống mã nguồn và cơ sở dữ liệu vận hành hoàn hảo, sẵn sàng cho buổi bảo vệ nghiệm thu trước Hội đồng Khoa Khoa học Máy tính.</li>
    </ul>

    <!-- CHỮ KÝ XÁC NHẬN -->
    <div class="footer-sign">
        <div>
            <div style="font-style: italic; margin-bottom: 3px;">Hà Nội, ngày 22 tháng 10 năm 2026</div>
            <div class="sign-title">GIẢNG VIÊN HƯỚNG DẪN</div>
            <div class="sign-name">ThS. Nguyễn Thị Thùy Liên<br><span style="font-size: 8.2pt; font-weight: normal; color: #64748b;">(Ký và ghi rõ họ tên)</span></div>
        </div>
        <div>
            <div style="font-style: italic; margin-bottom: 3px;">Hà Nội, ngày 22 tháng 10 năm 2026</div>
            <div class="sign-title">TẬP THỂ SINH VIÊN (NHÓM 4)</div>
            <div class="sign-name">Phạm Văn Đạt (Trưởng nhóm)<br>Nguyễn Duy Bảo &amp; Trần Tuấn Anh</div>
        </div>
    </div>

</body>
</html>
"""

# Ghi tệp HTML
with open(HTML_FILE, "w", encoding="utf-8") as f:
    f.write(html_content)
print(f"SUCCESS: Wrote HTML file at {HTML_FILE}")

# Xuất file PDF bằng Edge headless
if os.path.exists(EDGE_PATH):
    args = [
        EDGE_PATH,
        "--headless",
        "--disable-gpu",
        "--no-pdf-header-footer",
        f"--print-to-pdf={PDF_FILE}",
        f"file:///{HTML_FILE.replace(os.sep, '/')}"
    ]
    subprocess.run(args, check=True)
    print(f"SUCCESS: Generated PDF at {PDF_FILE}")

    try:
        shutil.copy2(PDF_FILE, DOWNLOAD_PDF)
        print(f"SUCCESS: Copied PDF to Downloads: {DOWNLOAD_PDF}")
    except Exception as e:
        print(f"Notice copy PDF to Downloads: {e}")

# ==============================================================================
# 2. TẠO TỆP WORD DOCX BÁO CÁO TIẾN ĐỘ CHUẨN FORMAL
# ==============================================================================
def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

doc = Document()

# Margin A4
for s in doc.sections:
    s.top_margin = Cm(2.0)
    s.bottom_margin = Cm(2.0)
    s.left_margin = Cm(2.2)
    s.right_margin = Cm(2.0)

# Default style
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(11)
font.color.rgb = RGBColor(0, 0, 0)
style.paragraph_format.line_spacing = 1.15
style.paragraph_format.space_after = Pt(3)

# Header
p_org = doc.add_paragraph()
p_org.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_org = p_org.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\nTRƯỜNG ĐẠI HỌC PHENIKAA – KHOA KHOA HỌC MÁY TÍNH\n***")
r_org.font.size = Pt(11)
r_org.bold = True

p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_before = Pt(8)
p_title.paragraph_format.space_after = Pt(4)
r_t = p_title.add_run("BÁO CÁO TIẾN ĐỘ DỰ ÁN\nHỆ THỐNG QUẢN LÝ VÀ DẪN ĐƯỜNG BÃI ĐỖ XE THÔNG MINH (SMART PARKING SYSTEM)")
r_t.font.size = Pt(14)
r_t.bold = True
r_t.font.color.rgb = RGBColor(30, 58, 138)

# Table info
t_info = doc.add_table(rows=4, cols=2)
t_info.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_info, color="CBD5E1", sz="4", val="single")
info_rows = [
    ("Học phần / Mã lớp:", "Phát triển ứng dụng Hệ thống thông tin (CSE702131-1-1-26(N01))"),
    ("Giảng viên hướng dẫn:", "ThS. Nguyễn Thị Thùy Liên"),
    ("Nhóm sinh viên thực hiện:", "Nhóm 4: Phạm Văn Đạt (Trưởng nhóm), Nguyễn Duy Bảo, Trần Tuấn Anh"),
    ("Đánh giá tiến độ hiện tại:", "HOÀN THÀNH 100% MỤC TIÊU ĐỀ TÀI – SẴN SÀNG NGHIỆM THU")
]
for idx, (k, v) in enumerate(info_rows):
    row = t_info.rows[idx]
    row.cells[0].width = Cm(5.5)
    row.cells[1].width = Cm(11.0)
    set_cell_margins(row.cells[0])
    set_cell_margins(row.cells[1])
    set_cell_background(row.cells[0], "F1F5F9")
    
    p0 = row.cells[0].paragraphs[0]
    p0.paragraph_format.space_after = Pt(0)
    r0 = p0.add_run(k)
    r0.bold = True
    r0.font.size = Pt(10.5)
    
    p1 = row.cells[1].paragraphs[0]
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run(v)
    r1.font.size = Pt(10.5)
    if idx == 3:
        r1.bold = True
        r1.font.color.rgb = RGBColor(22, 101, 52)

p_sp = doc.add_paragraph()
p_sp.paragraph_format.space_after = Pt(6)

def add_heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)

# I. BẢNG TIẾN ĐỘ TỪNG PHÂN HỆ
add_heading("I. BẢNG TIẾN ĐỘ CHI TIẾT TỪNG PHÂN HỆ NGHIỆP VỤ (WBS)")
headers_wbs = ["STT", "Phân hệ chức năng", "Nội dung kỹ thuật đã hoàn thành", "Phụ trách", "Tiến độ"]
data_wbs = [
    ["1", "Backend Core & CSDL", "6 bảng PostgreSQL 3NF, Alembic, Auth JWT/Bcrypt, WebSocket, Khóa SELECT FOR UPDATE", "Phạm Văn Đạt", "100%"],
    ["2", "Dẫn đường tối ưu A*", "Ma trận lưới 2D, hàm Heuristic Manhattan, Priority Queue heapq, phản hồi < 5ms", "Trần Tuấn Anh", "100%"],
    ["3", "AI Engine Dự báo lấp đầy", "Sinh 4.320 mẫu lịch sử, Random Forest 200 cây (MAE = 0.058), API dự báo chuỗi 24h", "Trần Tuấn Anh", "100%"],
    ["4", "Cổng Khách hàng (Customer)", "Giao diện Vue 3 Dark Cyber, sơ đồ lưới 2D, diễn hoạt xe GSAP 3, vé QR Base64, vé tháng", "Nguyễn Duy Bảo", "100%"],
    ["5", "Cổng Nhân viên (Staff Portal)", "Trực bãi phân công (assigned_lot), check-in quét QR & vãng lai, check-out mở barrier", "Đạt & Bảo", "100%"],
    ["6", "Cổng Quản trị viên (Admin)", "Quản lý bãi & slot, tạo nhân viên, duyệt vé tháng cư dân, báo cáo phân tích doanh thu", "Cả nhóm", "100%"]
]

t_wbs = doc.add_table(rows=len(data_wbs)+1, cols=5)
t_wbs.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_wbs, color="CBD5E1", sz="4", val="single")
col_w = [Cm(1.2), Cm(3.8), Cm(7.2), Cm(2.8), Cm(1.8)]

for c_idx, h in enumerate(headers_wbs):
    cell = t_wbs.rows[0].cells[c_idx]
    cell.width = col_w[c_idx]
    set_cell_background(cell, "E2E8F0")
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 4] else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(h)
    r.bold = True
    r.font.size = Pt(10)

for r_idx, row_vals in enumerate(data_wbs):
    row = t_wbs.rows[r_idx + 1]
    for c_idx, val in enumerate(row_vals):
        cell = row.cells[c_idx]
        cell.width = col_w[c_idx]
        set_cell_margins(cell)
        if r_idx % 2 == 1:
            set_cell_background(cell, "F8FAFC")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 4] else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(val)
        r.font.size = Pt(9.5)
        if c_idx == 4:
            r.bold = True
            r.font.color.rgb = RGBColor(22, 101, 52)

# II. KẾT QUẢ KIỂM THỬ TỰ ĐỘNG
add_heading("II. KẾT QUẢ KIỂM THỬ TỰ ĐỘNG & BẢO ĐẢM CHẤT LƯỢNG (QA)")
headers_qa = ["Bộ kiểm thử (Test Suite)", "Mục tiêu kiểm định", "Quy mô", "Kết quả"]
data_qa = [
    ["Unit Test Thuật toán A*", "test_astar.py: Né vật cản, điểm bắt đầu trùng đích, đường cụt", "6/6 cases", "100% Passed"],
    ["Unit Test AI Engine", "test_ml_engine.py: Phân phối chuẩn, tính tái lập, hội tụ mô hình", "5/5 cases", "100% Passed"],
    ["Concurrency Stress Test", "test_double_booking.py: 50 luồng đồng thời tranh chấp 1 slot", "50 workers", "Triệt tiêu xung đột (400)"],
    ["E2E Automation Testing", "Playwright: Luồng liên hoàn Khách đặt chỗ -> Staff quét QR mở Barrier", "6 kịch bản", "100% Ổn định"]
]
t_qa = doc.add_table(rows=len(data_qa)+1, cols=4)
t_qa.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_qa, color="CBD5E1", sz="4", val="single")
col_w_qa = [Cm(4.2), Cm(8.0), Cm(2.4), Cm(2.4)]

for c_idx, h in enumerate(headers_qa):
    cell = t_qa.rows[0].cells[c_idx]
    cell.width = col_w_qa[c_idx]
    set_cell_background(cell, "E2E8F0")
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [2, 3] else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(h)
    r.bold = True
    r.font.size = Pt(10)

for r_idx, row_vals in enumerate(data_qa):
    row = t_qa.rows[r_idx + 1]
    for c_idx, val in enumerate(row_vals):
        cell = row.cells[c_idx]
        cell.width = col_w_qa[c_idx]
        set_cell_margins(cell)
        if r_idx % 2 == 1:
            set_cell_background(cell, "F8FAFC")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [2, 3] else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(val)
        r.font.size = Pt(9.5)
        if c_idx == 3:
            r.bold = True
            r.font.color.rgb = RGBColor(22, 101, 52)

# III. KẾ HOẠCH BẢO VỆ
add_heading("III. HỒ SƠ DỰ ÁN VÀ KẾ HOẠCH NGHIỆM THU ĐỒ ÁN")
bullets = [
    "Báo cáo thuyết minh toàn diện hơn 85 trang chuẩn mẫu Phenikaa (BAO_CAO_DACS_Smart_Parking_System.docx) đã hoàn tất 100%.",
    "Bộ tài liệu phụ trợ đầy đủ: Kế hoạch 10 tuần, 2 Biên bản họp nhóm, Tuyển tập 50 câu hỏi & trả lời bảo vệ đồ án, Kịch bản thuyết trình và slide PowerPoint.",
    "Hệ thống phần mềm Full-Stack (Backend FastAPI, CSDL PostgreSQL, Frontend Vue 3 + Tailwind CSS + GSAP 3) khởi chạy mượt mà bằng 1 cú nhấp chuột qua file run_web.bat.",
    "Nhóm 4 đã sẵn sàng tự tin bảo vệ và nghiệm thu đề tài trước Hội đồng chấm thi Khoa Khoa học Máy tính."
]
for b in bullets:
    p_b = doc.add_paragraph()
    p_b.paragraph_format.left_indent = Cm(0.6)
    p_b.paragraph_format.space_after = Pt(2)
    r_sym = p_b.add_run("• ")
    r_sym.bold = True
    r_txt = p_b.add_run(b)
    r_txt.font.size = Pt(10.5)

# Chữ ký
p_sig_sp = doc.add_paragraph()
p_sig_sp.paragraph_format.space_before = Pt(16)

t_sign = doc.add_table(rows=1, cols=2)
t_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_sign, color="FFFFFF", sz="0", val="none")

c0 = t_sign.rows[0].cells[0]
c0.width = Cm(8.5)
p0 = c0.paragraphs[0]
p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
r0 = p0.add_run("Hà Nội, ngày 22 tháng 10 năm 2026\nGIẢNG VIÊN HƯỚNG DẪN\n\n\n\n\nThS. Nguyễn Thị Thùy Liên\n(Ký và ghi rõ họ tên)")
r0.font.size = Pt(10.5)
r0.bold = True

c1 = t_sign.rows[0].cells[1]
c1.width = Cm(8.5)
p1 = c1.paragraphs[0]
p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
r1 = p1.add_run("Hà Nội, ngày 22 tháng 10 năm 2026\nTẬP THỂ SINH VIÊN (NHÓM 4)\n\n\n\n\nPhạm Văn Đạt (Trưởng nhóm)\nNguyễn Duy Bảo & Trần Tuấn Anh")
r1.font.size = Pt(10.5)
r1.bold = True

# Lưu file DOCX
try:
    doc.save(DOCX_FILE)
    print(f"SUCCESS: Generated DOCX at {DOCX_FILE}")
except Exception as e:
    alt_docx = os.path.join(PROJECT_DIR, "Bao_cao_tien_do_du_an_Smart_Parking_System_Nhom4.docx")
    doc.save(alt_docx)
    DOCX_FILE = alt_docx
    print(f"SUCCESS: Generated DOCX at {DOCX_FILE}")

try:
    shutil.copy2(DOCX_FILE, DOWNLOAD_DOCX)
    print(f"SUCCESS: Copied DOCX to Downloads: {DOWNLOAD_DOCX}")
except Exception as e:
    print(f"Notice copy DOCX to Downloads: {e}")

print("=== HOÀN TẤT XUẤT BÁO CÁO TIẾN ĐỘ DỰ ÁN (CẢ PDF, DOCX VÀ HTML)! ===")

