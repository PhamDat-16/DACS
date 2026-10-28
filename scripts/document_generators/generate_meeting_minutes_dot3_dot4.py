import os
import sys
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
LOGO_PATH = os.path.join(ROOT_DIR, "extracted_media", "word", "media", "image1.png")
OUTPUT_DIR = os.path.join(ROOT_DIR, "bao_cao_de_xuat")
os.makedirs(OUTPUT_DIR, exist_ok=True)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="000000", sz="8", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_header_table(doc, topic_text="Đề tài: Smart Parking System – Hệ thống quản lý bãi đỗ xe thông minh"):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="000000", sz="8", val="single")

    widths = [Cm(3.8), Cm(9.6), Cm(4.2)]

    # Cột 1: Logo
    cell0 = table.rows[0].cells[0]
    cell0.width = widths[0]
    set_cell_margins(cell0, top=120, bottom=120, left=120, right=120)
    p0 = cell0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(LOGO_PATH):
        r_logo = p0.add_run()
        r_logo.add_picture(LOGO_PATH, width=Inches(1.2))
    else:
        r_txt = p0.add_run("PHENIKAA\nUNIVERSITY")
        r_txt.font.name = 'Times New Roman'
        r_txt.font.size = Pt(10)
        r_txt.bold = True

    # Cột 2: Tiêu đề biên bản
    cell1 = table.rows[0].cells[1]
    cell1.width = widths[1]
    set_cell_margins(cell1, top=120, bottom=120, left=120, right=120)
    p1 = cell1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.line_spacing = 1.15
    p1.paragraph_format.space_before = Pt(2)
    p1.paragraph_format.space_after = Pt(2)

    r1_1 = p1.add_run("BIÊN BẢN HỌP NHÓM VÀ BÁO CÁO TIẾN ĐỘ CÔNG VIỆC\n")
    r1_1.font.name = 'Times New Roman'
    r1_1.font.size = Pt(11)
    r1_1.bold = True

    r1_2 = p1.add_run("(Học phần: Phát triển ứng dụng Hệ thống thông tin)\n")
    r1_2.font.name = 'Times New Roman'
    r1_2.font.size = Pt(9.5)
    r1_2.italic = True

    r1_3 = p1.add_run(topic_text)
    r1_3.font.name = 'Times New Roman'
    r1_3.font.size = Pt(10)
    r1_3.bold = True

    # Cột 3: Mã số BM
    cell2 = table.rows[0].cells[2]
    cell2.width = widths[2]
    set_cell_margins(cell2, top=120, bottom=120, left=120, right=120)
    p2 = cell2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p2.paragraph_format.line_spacing = 1.15
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(2)

    r2_1 = p2.add_run("Mã số: BM.CNTT.01.06\nLần ban hành: 01")
    r2_1.font.name = 'Times New Roman'
    r2_1.font.size = Pt(9.5)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(6)

def add_signature_table(doc, leader_name="Phạm Văn Đạt", instructor_name="ThS. Nguyễn Thị Thùy Liên"):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="FFFFFF", sz="0", val="none")

    widths = [Cm(8.8), Cm(8.8)]

    # Cột 1: Giảng viên
    c0 = table.rows[0].cells[0]
    c0.width = widths[0]
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.line_spacing = 1.15
    r0 = p0.add_run(f"GIẢNG VIÊN HƯỚNG DẪN\n\n\n\n\n{instructor_name}\n(Ký và ghi rõ họ tên)")
    r0.font.name = 'Times New Roman'
    r0.font.size = Pt(11)
    r0.bold = True

    # Cột 2: Nhóm trưởng
    c1 = table.rows[0].cells[1]
    c1.width = widths[1]
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.line_spacing = 1.15
    r1 = p1.add_run(f"NHÓM TRƯỞNG\n\n\n\n\n{leader_name}\n(Ký và ghi rõ họ tên)")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11)
    r1.bold = True

def add_heading_section(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.bold = True
    return p

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.6)
    r_sym = p.add_run("• ")
    r_sym.font.name = 'Times New Roman'
    r_sym.font.size = Pt(10.5)
    r_sym.bold = True
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(10.5)
        r_pre.bold = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10.5)
    return p

def create_meeting_doc(filename, time_str, period_str, details, tasks, plans, conclusion_text):
    print(f"-> Đang tạo tài liệu: {filename} ...")
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.70)
        s.bottom_margin = Cm(1.70)
        s.left_margin = Cm(1.80)
        s.right_margin = Cm(1.80)
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)

    add_header_table(doc, "Đề tài: Smart Parking System – Hệ thống quản lý bãi đỗ xe thông minh")

    add_heading_section(doc, "1. Thời gian và địa điểm")
    add_bullet(doc, time_str, bold_prefix="Thời gian: ")
    add_bullet(doc, "Họp trực tuyến qua ứng dụng Google Meet kết hợp kiểm thử trực tiếp trên môi trường máy chủ cục bộ.", bold_prefix="Địa điểm: ")

    add_heading_section(doc, "2. Thành phần tham dự")
    add_bullet(doc, "Đầy đủ 03 thành viên thuộc Nhóm 4 nghiên cứu đề tài:")
    add_bullet(doc, "Phạm Văn Đạt – Trưởng nhóm / Chủ tọa cuộc họp (MSV: 23017150).")
    add_bullet(doc, "Nguyễn Duy Bảo – Kỹ sư Giao diện / Thư ký cuộc họp (MSV: 23017133).")
    add_bullet(doc, "Trần Tuấn Anh – Kỹ sư Thuật toán & AI / Thành viên (MSV: 23017118).")

    add_heading_section(doc, "3. Nội dung chi tiết")
    for item in details:
        add_bullet(doc, item)

    add_heading_section(doc, "3.1. Thời gian thực hiện")
    add_bullet(doc, period_str)

    add_heading_section(doc, "3.2. Đánh giá tiến độ công việc")

    headers_t = ["TT", "Mã sinh viên", "Họ và tên", "Nhiệm vụ phân công", "Nội dung đã/đang thực hiện", "Đánh giá tiến độ", "Ghi chú"]
    t = doc.add_table(rows=len(tasks) + 1, cols=len(headers_t))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t, color="000000", sz="8", val="single")

    col_w = [Cm(1.0), Cm(2.2), Cm(3.0), Cm(3.8), Cm(4.6), Cm(2.2), Cm(2.4)]

    for c_idx, heading in enumerate(headers_t):
        cell = t.rows[0].cells[c_idx]
        cell.width = col_w[c_idx]
        set_cell_background(cell, "E7E6E6")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(heading)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.5)
        r.bold = True

    for r_idx, r_vals in enumerate(tasks):
        row = t.rows[r_idx + 1]
        for c_idx, val in enumerate(r_vals):
            cell = row.cells[c_idx]
            cell.width = col_w[c_idx]
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            if c_idx in [0, 1, 5]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9.0)

    p_sp_t = doc.add_paragraph()
    p_sp_t.paragraph_format.space_after = Pt(4)

    add_heading_section(doc, "4. Kế hoạch công việc sau cuộc họp")
    for p_item in plans:
        add_bullet(doc, p_item)

    add_heading_section(doc, "5. Kết luận cuộc họp")
    p_conclusion = doc.add_paragraph()
    p_conclusion.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_conclusion.paragraph_format.space_before = Pt(0)
    p_conclusion.paragraph_format.space_after = Pt(8)
    p_conclusion.paragraph_format.line_spacing = 1.15
    r_c = p_conclusion.add_run(conclusion_text)
    r_c.font.name = 'Times New Roman'
    r_c.font.size = Pt(10.5)

    p_sp_sig = doc.add_paragraph()
    p_sp_sig.paragraph_format.space_after = Pt(8)

    add_signature_table(doc, leader_name="Phạm Văn Đạt")

    full_save_path = os.path.join(OUTPUT_DIR, filename)
    doc.save(full_save_path)
    print(f"==> Đã lưu file tại: {full_save_path}")

    # Đồng thời sao chép sang Downloads
    user_dl = os.path.join(os.path.expanduser("~"), "Downloads")
    if os.path.exists(user_dl):
        dl_target = os.path.join(user_dl, filename)
        try:
            shutil.copy2(full_save_path, dl_target)
            print(f"==> Đã sao chép sang Downloads: {dl_target}")
        except Exception as e:
            print(f"Lưu ý: Không thể sao chép sang Downloads: {e}")

# ==============================================================================
# 1. BIÊN BẢN HỌP NHÓM ĐỢT 3 (LẦN 3) - Ngày 14/10/2026
# ==============================================================================
def generate_dot_3():
    filename = "Bien_ban_hop_nhom_SmartParking_Dot_3_14-10-2026_Nhom4.docx"
    time_str = "14h30 – 17h00, ngày 14/10/2026."
    period_str = "Giai đoạn Đợt 3 (Tuần 7 – Tuần 8): từ 05/10/2026 đến 18/10/2026."

    details = [
        "Đánh giá tiến độ triển khai Đợt 3 của đồ án Smart Parking System: Hoàn thiện cổng kiểm soát dành cho Nhân viên vận hành (Staff Portal) với cơ chế phân quyền bãi trực nghiêm ngặt (assigned_lot).",
        "Tích hợp kết nối WebSocket hai chiều (app/api/websocket.py) cập nhật trạng thái ô đỗ thời gian thực: khi khách đặt chỗ hoặc check-in, ma trận ô đỗ lập tức đổi màu (Xanh: Trống; Đỏ: Có xe; Vàng: Đã đặt trước) tức thì trên toàn bộ các trình duyệt đang mở.",
        "Tích hợp thành công thư viện quét mã QR Code trực tiếp qua camera/webcam (html5-qrcode) tại cổng barrier: hệ thống tự động giải mã chuỗi JWT/Token vé xe, kiểm tra tính hợp lệ và kích hoạt hoạt ảnh mở cần barrier tự động.",
        "Xây dựng và kiểm thử luồng Check-in xe vãng lai (Walk-in check-in): nhân viên chỉ cần nhấp một nút, hệ thống tự động tìm và cấp phát ô đỗ trống gần cổng nhất cho khách không đặt trước.",
        "Hiện thực hóa và huấn luyện mô hình Trí tuệ nhân tạo (Random Forest Regressor): Xây dựng pipeline xử lý 4.320 bản ghi lịch sử lấp đầy theo 24 khung giờ x 180 ngày; mô hình đạt 200 Decision Trees, sai số MAE = 0.058 và được xuất bản thành file Singleton model (.joblib) nạp sẵn trong bộ nhớ RAM phục vụ dự báo tỷ lệ lấp đầy trong 24 giờ tới.",
        "Tiếp tục chuẩn hóa cơ sở dữ liệu PostgreSQL và kiểm thử khóa giao dịch SELECT FOR UPDATE, đảm bảo khả năng chịu tải khi nhiều yêu cầu đặt chỗ gửi đến cùng một thời điểm."
    ]

    tasks = [
        [
            "1",
            "23017150",
            "Phạm Văn Đạt",
            "WebSocket ConnectionManager – Staff Portal API – SELECT FOR UPDATE",
            "Hoàn thành kiến trúc phát sóng WebSocket real-time; viết API phân quyền nhân viên theo assigned_lot; tích hợp logic xác thực vé QR và mở barrier tự động.",
            "Đúng kế hoạch",
            "Tiếp tục hoàn thiện luồng thanh toán Check-out và tính phí gửi xe theo giờ."
        ],
        [
            "2",
            "23017133",
            "Nguyễn Duy Bảo",
            "Staff Portal UI – Webcam QR Scanner – Barrier Animation Modal",
            "Xây dựng giao diện bảng điều khiển nhân viên trực bãi; tích hợp thư viện html5-qrcode quét vé qua webcam; thiết kế hiệu ứng cần barrier nâng hạ bằng CSS/GSAP.",
            "Đúng kế hoạch",
            "Phối hợp kiểm thử thực tế quét vé QR trên nhiều loại điện thoại và góc camera."
        ],
        [
            "3",
            "23017118",
            "Trần Tuấn Anh",
            "Pipeline Dữ liệu 4.320 mẫu – Random Forest Model – AI Inference API",
            "Hoàn thành script tiền xử lý 4.320 mẫu dữ liệu; huấn luyện Random Forest Regressor 200 cây đạt MAE = 0.058; đóng gói file .joblib và tích hợp endpoint /api/predict.",
            "Đúng kế hoạch",
            "Chuẩn bị các kịch bản kiểm thử tải đồng thời (Concurrency Stress Test)."
        ]
    ]

    plans = [
        "Hoàn thiện luồng Check-out tự động: quét mã vé lúc xe rời bãi, tự động tính tổng thời gian đỗ, áp dụng biểu phí theo giờ và xuất biên lai thanh toán.",
        "Nghiên cứu và triển khai tính năng Vé tháng điện tử (Monthly Pass) dành cho cư dân và nhân viên văn phòng.",
        "Mở rộng bộ kiểm thử tự động Pytest bao phủ các endpoint của Staff Portal, WebSocket và Random Forest Engine.",
        "Mỗi thành viên tiếp tục kiểm thử đơn vị trên branch cá nhân trước khi tạo Pull Request hợp nhất vào nhánh develop.",
        "Chụp ảnh màn hình các tính năng mới hoàn thành để đưa vào Báo cáo tiến độ và Sổ tay hướng dẫn sử dụng."
    ]

    conclusion = (
        "Cuộc họp Đợt 3 kết thúc thành công, các mục tiêu cốt lõi của giai đoạn Tuần 7 – Tuần 8 đã hoàn thành xuất sắc 100%. "
        "Hệ thống đã vận hành thông suốt từ khâu Khách hàng đặt chỗ trực tuyến, dẫn đường A*, truyền tin thời gian thực qua WebSocket "
        "đến khâu Nhân viên quét mã QR mở barrier tại cổng và dự báo lấp đầy bằng AI. Toàn đội thống nhất dồn lực cho Đợt 4 nhằm hoàn thiện "
        "tính năng Vé tháng, hoàn thành 24/24 ca kiểm thử Pytest và đóng gói toàn bộ hồ sơ nghiệm thu đồ án."
    )

    create_meeting_doc(filename, time_str, period_str, details, tasks, plans, conclusion)

# ==============================================================================
# 2. BIÊN BẢN HỌP NHÓM ĐỢT 4 (LẦN 4) - Ngày 28/10/2026
# ==============================================================================
def generate_dot_4():
    filename = "Bien_ban_hop_nhom_SmartParking_Dot_4_28-10-2026_Nhom4.docx"
    time_str = "19h00 – 21h30, ngày 28/10/2026."
    period_str = "Giai đoạn Đợt 4 (Tuần 9 – Tuần 10): từ 19/10/2026 đến 30/10/2026 (Hoàn thành 100% đồ án)."

    details = [
        "Đánh giá tổng kết toàn diện giai đoạn nước rút Đợt 4 (Tuần 9 – Tuần 10) của đồ án Smart Parking System: Hệ thống đã hoàn thiện 100% các tính năng đề ra, đạt trạng thái sẵn sàng nghiệm thu và bảo vệ trước Hội đồng chấm điểm.",
        "Nghiệm thu tính năng Vé tháng điện tử (Monthly Pass) và Quản lý phương tiện cá nhân: Cư dân đăng ký vé tháng trực tuyến (gói 30 ngày, 90 ngày), xe vé tháng khi đến cổng barrier được hệ thống nhận diện tức thì và miễn phí gửi; bổ sung modal tra cứu lịch sử đỗ xe chi tiết cho từng phương tiện trong giao diện Quản trị viên.",
        "Báo cáo kết quả kiểm thử tự động toàn diện (Automated Testing QA): Chạy bộ kiểm thử tự động Pytest đạt tuyệt đối 24/24 ca kiểm thử PASSED 100% (kiểm thử giải thuật A*, suy luận AI Random Forest, Stress Test 50 workers chống Race Condition khóa SELECT FOR UPDATE, kiểm tra chu kỳ vé tháng và đồng bộ lịch sử xe).",
        "Khắc phục triệt để lỗi kiểm thử vé tháng tĩnh bằng cơ chế Dynamic Date Isolation trong Pytest fixture, đảm bảo bộ kiểm thử luôn chạy ổn định vĩnh viễn theo thời gian.",
        "Đóng gói công cụ khởi chạy tự động 1-Click (run.bat, run_web.bat, xem_database.bat) và xuất bản bản sao lưu cơ sở dữ liệu PostgreSQL hoàn chỉnh (smart_parking_backup.sql), cho phép khởi động toàn bộ hệ sinh thái chỉ bằng 1 cú nhấp chuột.",
        "Hoàn thiện trọn bộ hồ sơ đồ án phục vụ bảo vệ: Báo cáo toàn văn Đồ án tốt nghiệp 100 trang chuẩn mẫu Đại học Phenikaa (Bao_cao_do_an_Smart_Parking_System.docx), Bộ 8 biểu đồ Use Case chuẩn UML, Tuyển tập 50 câu hỏi & đáp án bảo vệ đồ án, Kịch bản thuyết trình và Slide trình chiếu."
    ]

    tasks = [
        [
            "1",
            "23017150",
            "Phạm Văn Đạt",
            "Monthly Pass API – Vehicle History Sync – 24 Pytest Cases – 1-Click Bat",
            "Hoàn thành trọn gói tính năng vé tháng VIP và đồng bộ lịch sử đỗ xe; tối ưu hóa 24/24 test cases Pytest; xây dựng file run_web.bat và xem_database.bat; đóng gói smart_parking_backup.sql.",
            "Hoàn thành xuất sắc",
            "Sẵn sàng demo trực tiếp toàn bộ luồng nghiệp vụ trước Hội đồng."
        ],
        [
            "2",
            "23017133",
            "Nguyễn Duy Bảo",
            "Vehicle History Modal – Search Filter – Admin Reports – Presentation Slide",
            "Xây dựng modal hiển thị lịch sử gửi đỗ của từng phương tiện; nâng cấp bộ tìm kiếm đa năng theo biển số / mã booking / QR token; thiết kế bộ Slide trình chiếu thuyết minh đồ án chuyên nghiệp.",
            "Hoàn thành xuất sắc",
            "Chuẩn bị máy chiếu và điều khiển trình chiếu trong buổi bảo vệ."
        ],
        [
            "3",
            "23017118",
            "Trần Tuấn Anh",
            "Concurrency Benchmark – A* Optimization – Full Thesis Documentation",
            "Thực hiện đo kiểm tải đồng thời chứng minh hiệu quả khóa bi quan SELECT FOR UPDATE; tối ưu hóa tốc độ A* dưới 3.5ms; biên soạn và rà soát kỹ thuật toàn bộ 100 trang Báo cáo Đồ án tốt nghiệp.",
            "Hoàn thành xuất sắc",
            "Phụ trách phần thuyết minh giải thuật A* và mô hình học máy AI."
        ]
    ]

    plans = [
        "In ấn và đóng bìa mạ vàng toàn bộ tập Báo cáo Đồ án tốt nghiệp theo quy định của Khoa Công nghệ thông tin – Đại học Phenikaa.",
        "Tiến hành tổng duyệt buổi thuyết trình thử (Mock Defense): chạy thử nghiệm slide, phân chia thời lượng trình bày cho từng thành viên (15 phút thuyết trình + 10 phút trả lời phản biện).",
        "Kiểm tra môi trường demo trên máy tính xách tay: bảo đảm PostgreSQL, Backend FastAPI và Frontend Web chạy độc lập ngoại tuyến (offline) mượt mà không phụ thuộc Internet.",
        "Chuẩn bị đầy đủ phụ lục hồ sơ minh chứng: Kế hoạch 10 tuần, Trọn bộ 4 Biên bản họp nhóm (Đợt 1, 2, 3, 4), Bộ 8 biểu đồ Use Case và Kết quả đo kiểm hiệu năng."
    ]

    conclusion = (
        "Cuộc họp Đợt 4 – phiên họp tổng kết đề tài Smart Parking System – đã hoàn thành xuất sắc toàn bộ nội dung. "
        "Sau 10 tuần nỗ lực làm việc nghiêm túc, nhóm nghiên cứu đã xây dựng thành công một giải pháp quản lý bãi đỗ xe thông minh toàn diện, "
        "kết hợp hài hòa giữa kỹ thuật phần mềm hiện đại (FastAPI, Vue 3, WebSocket), thuật toán tìm đường tối ưu A*, mô hình học máy AI Random Forest "
        "và cơ chế an toàn dữ liệu cao cấp. Nhóm 4 tự tin khẳng định đề tài đã sẵn sàng 100% để báo cáo và bảo vệ đạt kết quả cao nhất trước Hội đồng chấm điểm."
    )

    create_meeting_doc(filename, time_str, period_str, details, tasks, plans, conclusion)

if __name__ == "__main__":
    generate_dot_3()
    generate_dot_4()
    print("\n===> TẤT CẢ 2 BIÊN BẢN HỌP NHÓM ĐỢT 3 VÀ 4 ĐÃ ĐƯỢC TẠO THÀNH CÔNG!")
