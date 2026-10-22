import os
import sys
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

if sys.stdout.encoding != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="000000", sz="8", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def build_smart_parking_plan():
    doc = Document()
    
    # 1. Cấu hình lề trang giống hệt BloomGift
    for s in doc.sections:
        s.top_margin = Cm(1.70)
        s.bottom_margin = Cm(1.70)
        s.left_margin = Cm(1.80)
        s.right_margin = Cm(1.80)
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)

    # Cấu hình Default Style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 1.08
    style.paragraph_format.space_after = Pt(2)
    style.paragraph_format.space_before = Pt(0)

    # Tiêu đề chính 1
    p0 = doc.add_paragraph()
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_after = Pt(2)
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.line_spacing = 1.08
    r0 = p0.add_run("KẾ HOẠCH THỰC HIỆN ĐỒ ÁN HỆ THỐNG THÔNG TIN")
    r0.font.name = 'Times New Roman'
    r0.font.size = Pt(15)
    r0.bold = True
    r0.font.color.rgb = RGBColor(0, 0, 0)

    # Tiêu đề chính 2
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_after = Pt(8)
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.line_spacing = 1.08
    r1 = p1.add_run("ĐỀ TÀI: SMART PARKING SYSTEM – HỆ THỐNG QUẢN LÝ BÃI ĐỖ XE THÔNG MINH")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(13)
    r1.bold = True
    r1.font.color.rgb = RGBColor(0, 0, 0)

    # Dòng trống nhẹ
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(2)

    # Bảng 1: Bảng thông tin thực hiện đồ án (Table 1)
    t1_data = [
        [("Thời gian thực hiện", True), ("26/08/2026 – 30/10/2026", False), ("Số thành viên", True), ("03 thành viên (Nhóm 4)", False)],
        [("Thành viên 1", True), ("Phạm Văn Đạt (Trưởng nhóm)", False), ("Thành viên 2", True), ("Nguyễn Duy Bảo (Thành viên)", False)],
        [("Thành viên 3", True), ("Trần Tuấn Anh (Thành viên)", False), ("Tổng số tuần", True), ("10 tuần", False)],
        [("Mô hình thực hiện", True), ("Mô hình Agile / Scrum (Lặp và tăng trưởng)", False), ("Giảng viên hướng dẫn", True), ("ThS. Nguyễn Thị Thùy Liên", False)]
    ]

    table1 = doc.add_table(rows=4, cols=4)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table1, color="000000", sz="8", val="single")

    col_widths_t1 = [Cm(3.8), Cm(4.9), Cm(3.8), Cm(4.9)]

    for r_idx, row in enumerate(t1_data):
        table_row = table1.rows[r_idx]
        for c_idx, (text, is_label) in enumerate(row):
            cell = table_row.cells[c_idx]
            cell.width = col_widths_t1[c_idx]
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            if is_label:
                set_cell_background(cell, "E7E6E6")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10.5)
            r.bold = is_label
            r.font.color.rgb = RGBColor(0, 0, 0)

    # Dòng trống phân cách
    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_after = Pt(4)

    # Phần 1. MỤC TIÊU KẾ HOẠCH
    p_sec1 = doc.add_paragraph()
    p_sec1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_sec1.paragraph_format.space_before = Pt(4)
    p_sec1.paragraph_format.space_after = Pt(4)
    p_sec1.paragraph_format.line_spacing = 1.08
    r_s1 = p_sec1.add_run("1. MỤC TIÊU KẾ HOẠCH")
    r_s1.font.name = 'Times New Roman'
    r_s1.font.size = Pt(12.5)
    r_s1.bold = True
    r_s1.font.color.rgb = RGBColor(0, 0, 0)

    p_body1 = doc.add_paragraph()
    p_body1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_body1.paragraph_format.space_before = Pt(0)
    p_body1.paragraph_format.space_after = Pt(6)
    p_body1.paragraph_format.line_spacing = 1.08
    r_b1 = p_body1.add_run("Kế hoạch được xây dựng theo đúng trình tự của một dự án phần mềm chuyên nghiệp: khảo sát nghiệp vụ và phân tích yêu cầu trước, thiết kế kiến trúc hệ thống và cơ sở dữ liệu sau đó mới triển khai lập trình backend/frontend, tích hợp thuật toán điều hướng A* và mô hình học máy Random Forest, kiểm thử đa luồng chống xung đột và hoàn thiện sản phẩm. Mỗi tuần đều có nội dung thực hiện cụ thể và kết quả mong muốn để phục vụ báo cáo tiến độ.")
    r_b1.font.name = 'Times New Roman'
    r_b1.font.size = Pt(11.0)
    r_b1.font.color.rgb = RGBColor(0, 0, 0)

    # Phần 2. KẾ HOẠCH THỰC HIỆN THEO TỪNG TUẦN
    p_sec2 = doc.add_paragraph()
    p_sec2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_sec2.paragraph_format.space_before = Pt(6)
    p_sec2.paragraph_format.space_after = Pt(4)
    p_sec2.paragraph_format.line_spacing = 1.08
    r_s2 = p_sec2.add_run("2. KẾ HOẠCH THỰC HIỆN THEO TỪNG TUẦN")
    r_s2.font.name = 'Times New Roman'
    r_s2.font.size = Pt(12.5)
    r_s2.bold = True
    r_s2.font.color.rgb = RGBColor(0, 0, 0)

    # Bảng 2: Kế hoạch 10 tuần chi tiết
    t2_headers = ["Tuần", "Thời gian", "Kế hoạch thực hiện", "Mong muốn đạt được", "Phụ trách chính"]
    
    t2_data = [
        [
            "Tuần 1",
            "26/08 – 30/08/2026",
            "Khởi động đề tài Smart Parking System (SPS); khảo sát hiện trạng các bãi giữ xe truyền thống và bán tự động tại Việt Nam; xác định mục tiêu, phạm vi nghiên cứu; phân định 3 nhóm tác nhân Customer, Staff và Admin; thống nhất danh mục tính năng cốt lõi (Bản đồ 2D, Dẫn đường A*, Đặt chỗ trực tuyến, Quét QR Barrier, Dự báo AI).",
            "Xác định rõ đề tài, mục tiêu, phạm vi và các tác nhân. Hoàn thành đề xuất dự án (Project Proposal) và danh mục chức năng làm cơ sở cho bước phân tích nghiệp vụ.",
            "Cả nhóm (Đạt, Bảo, Tuấn Anh)"
        ],
        [
            "Tuần 2",
            "31/08 – 06/09/2026",
            "Phân tích BA và yêu cầu hệ thống; xác định yêu cầu chức năng, yêu cầu phi chức năng (bảo mật JWT, độ trễ WebSocket < 50ms, chống xung đột vé) và quy tắc nghiệp vụ; phân tích các luồng Đăng ký, Đăng nhập, Tra cứu bãi xe, Đặt giữ chỗ, Dẫn đường A*, Quét QR Check-in/out, Đăng ký vé tháng và Dự báo lấp đầy; xây dựng sơ đồ Use Case tổng quát và đặc tả chi tiết 13 Use Cases.",
            "Hoàn thành tài liệu phân tích yêu cầu phần mềm (SRS). Các tác nhân và luồng nghiệp vụ đỗ xe được chuẩn hóa rõ ràng trước khi thiết kế kỹ thuật.",
            "Cả nhóm (Đạt: SRS/Backend; Bảo: Use Case/UI; Tuấn Anh: Thuật toán/AI)"
        ],
        [
            "Tuần 3",
            "07/09 – 13/09/2026",
            "Thiết kế cơ sở dữ liệu quan hệ và kiến trúc hệ thống 3 tầng; xây dựng sơ đồ ERD chuẩn hóa 3NF; đặc tả chi tiết 6 bảng thực thể: users, parking_lots, parking_slots, bookings, sessions (hoặc monthly_passes) và occupancy_history; thiết kế ma trận lưới bãi đỗ 2D (Grid Matrix); phác thảo Wireframe giao diện Dashboard và sơ đồ điều hướng Sitemap.",
            "Có mô hình dữ liệu quan hệ hoàn chỉnh, quan hệ khóa chính – khóa ngoại rõ ràng, kiến trúc Client-Server (FastAPI + Vue 3) và bản vẽ wireframe các màn hình chính.",
            "Đạt: CSDL/ERD/Arch; Bảo: Wireframe/UI; Tuấn Anh: Lưới tọa độ"
        ],
        [
            "Tuần 4",
            "14/09 – 20/09/2026",
            "Hoàn thiện thiết kế chi tiết bằng Activity Diagram, Sequence Diagram cho các luồng then chốt: Đặt chỗ (Booking Flow), Check-in/Check-out qua Barrier và Dẫn đường A*; nghiên cứu giải pháp kỹ thuật SELECT FOR UPDATE chống Race Condition; nghiên cứu cơ sở lý thuyết toán học thuật toán A* (hàm f = g + h, khoảng cách Manhattan) và mô hình hồi quy Random Forest; đặc tả 28 REST API Endpoints.",
            "Hoàn thành bộ tài liệu phân tích – thiết kế kiến trúc chi tiết, giải pháp thuật toán và thiết kế API làm nền tảng vững chắc cho việc triển khai mã nguồn.",
            "Đạt: API Spec & Khóa hàng; Bảo: Sequence UI; Tuấn Anh: Toán học A* & RF"
        ],
        [
            "Tuần 5",
            "21/09 – 27/09/2026",
            "Khởi tạo và cấu hình môi trường dự án; thiết lập Backend FastAPI với SQLAlchemy ORM, Alembic migration và PostgreSQL; cài đặt module Authentication (băm mật khẩu Bcrypt, cấp phát và giải mã JWT); tạo seed data khởi tạo tài khoản Admin, Staff và các bãi xe mẫu; khởi tạo Frontend Vue 3 (Vite, Tailwind CSS, Pinia, GSAP 3).",
            "Backend và Frontend khởi chạy thành công, kết nối PostgreSQL ổn định, các API xác thực danh tính (Register/Login) và phân quyền Role hoạt động hoàn chỉnh.",
            "Đạt: Backend Core/Auth; Bảo: Vue 3 Setup/Tailwind; Tuấn Anh: Khởi tạo mô hình AI"
        ],
        [
            "Tuần 6",
            "28/09 – 04/10/2026",
            "Hiện thực hóa thuật toán A* (app/pathfinding/astar.py) với hàng đợi heapq và hàm Heuristic Manhattan; xây dựng component bản đồ ParkingGrid.vue và hiệu ứng xe di chuyển mượt mà bám lộ trình bằng GSAP 3; hiện thực hóa API Đặt chỗ (POST /api/bookings) tích hợp mệnh đề SELECT FOR UPDATE; sinh mã vé điện tử QR Code động cho khách hàng.",
            "Khách hàng xem được bản đồ bãi xe, bấm chọn slot trống được thuật toán vẽ đường chỉ dẫn xe chạy, đặt chỗ thành công có mã QR vé xe và không bị trùng lặp ô đỗ.",
            "Đạt: Booking API/Khóa hàng; Bảo: GSAP Animation/Grid UI; Tuấn Anh: Thuật toán A* Heapq"
        ],
        [
            "Tuần 7",
            "05/10 – 11/10/2026",
            "Hiện thực hóa Cổng nghiệp vụ Nhân viên (Staff Portal) và Quản trị viên (Admin Portal): Xây dựng tính năng phân quyền nhân viên theo bãi làm việc (assigned_lot); kết nối giao tiếp thời gian thực WebSocket (cập nhật trạng thái slot xanh/đỏ tức thì); tích hợp trình quét mã QR bằng webcam/camera và cửa sổ điều khiển mở Barrier tự động; lập trình API tính phí đỗ theo giờ và in hóa đơn điện tử.",
            "Nhân viên quản lý đúng bãi xe được phân công; quét mã QR mở barrier thành công; trạng thái ô đỗ trên lưới thay đổi thời gian thực qua WebSocket; tính phí tự động khi check-out.",
            "Đạt: WebSocket/Barrier Gate; Bảo: Staff Portal/Quét QR; Tuấn Anh: Pipeline dữ liệu AI"
        ],
        [
            "Tuần 8",
            "12/10 – 18/10/2026",
            "Xây dựng Machine Learning Engine dự báo tỷ lệ lấp đầy bãi xe: Viết pipeline sinh 4.320 mẫu dữ liệu lịch sử (180 ngày x 24 giờ); tiền xử lý 8 đặc trưng đầu vào; huấn luyện mô hình Random Forest Regressor với 200 Decision Trees; đánh giá độ chính xác (MAE ~ 0.058, R² ~ 0.912); đóng gói file mô hình .joblib và tích hợp nạp Singleton vào API FastAPI; vẽ biểu đồ dự báo 24h trên Admin Dashboard.",
            "Mô hình AI dự báo chính xác cao; API dự báo phản hồi dưới 5ms; quản trị viên theo dõi được biểu đồ xu hướng lấp đầy 24 giờ tới để điều phối hoạt động bãi xe.",
            "Tuấn Anh: Huấn luyện Random Forest; Đạt: API AI/Joblib; Bảo: Chart.js Admin UI"
        ],
        [
            "Tuần 9",
            "19/10 – 25/10/2026",
            "Tích hợp toàn diện hệ thống Client-Server; xây dựng bộ kiểm thử tự động Pytest (11 unit tests đạt 100% Passed); thực hiện bài test kiểm thử tương tranh đa luồng (50 workers) chứng minh triệt tiêu lỗi Double-booking; kiểm thử tự động hóa giao diện E2E bằng Playwright; kiểm thử các trường hợp biên, tối ưu hóa hiệu năng API (< 100ms) và vá lỗi giao diện.",
            "Hệ thống vận hành trơn tru, ổn định tuyệt đối dưới tải cao; bộ kiểm thử Pytest và Playwright đầy đủ minh chứng Đạt/Không đạt; tài liệu kiểm thử hoàn tất.",
            "Đạt: Pytest/Test đa luồng; Bảo: Playwright E2E; Tuấn Anh: Đánh giá sai số MAE/RMSE"
        ],
        [
            "Tuần 10",
            "26/10 – 30/10/2026",
            "Rà soát mã nguồn, kiểm tra lại toàn bộ tài khoản Admin/Staff; cập nhật README.md hướng dẫn cài đặt và khởi chạy 1-click (run_web.bat); hoàn thiện Báo cáo đồ án đầy đủ chuẩn mẫu Phenikaa (hơn 90 trang .docx); tổng hợp Tuyển tập 50 câu hỏi & trả lời bảo vệ đồ án; chuẩn bị slide thuyết trình và kịch bản demo hệ thống phục vụ hội đồng nghiệm thu.",
            "Có bộ hồ sơ đồ án xuất sắc gồm source code, database, báo cáo Word chi tiết, tài liệu 50 câu hỏi bảo vệ, slide trình chiếu; hệ thống sẵn sàng nghiệm thu và bảo vệ đạt điểm tối đa.",
            "Cả nhóm (Đạt: Báo cáo kỹ thuật; Bảo: Slide/Demo; Tuấn Anh: Bộ câu hỏi bảo vệ)"
        ]
    ]

    table2 = doc.add_table(rows=len(t2_data) + 1, cols=len(t2_headers))
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table2, color="000000", sz="8", val="single")

    col_widths_t2 = [Cm(1.4), Cm(2.5), Cm(6.5), Cm(4.8), Cm(2.2)]

    # Dòng Header
    hdr_row = table2.rows[0]
    for c_idx, heading in enumerate(t2_headers):
        cell = hdr_row.cells[c_idx]
        cell.width = col_widths_t2[c_idx]
        set_cell_background(cell, "D9D9D9")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(heading)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.5)
        r.bold = True
        r.font.color.rgb = RGBColor(0, 0, 0)

    # Các dòng dữ liệu tuần 1 đến tuần 10
    for r_idx, row_vals in enumerate(t2_data):
        table_row = table2.rows[r_idx + 1]
        for c_idx, val in enumerate(row_vals):
            cell = table_row.cells[c_idx]
            cell.width = col_widths_t2[c_idx]
            set_cell_margins(cell, top=90, bottom=90, left=110, right=110)
            p = cell.paragraphs[0]
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif c_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif c_idx == 4:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05

            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9.0)
            if c_idx == 0:
                r.bold = True
            r.font.color.rgb = RGBColor(0, 0, 0)

    # Dòng trống phân cách
    p_sp3 = doc.add_paragraph()
    p_sp3.paragraph_format.space_after = Pt(4)

    # Phần 3. KẾT QUẢ MONG MUỐN CUỐI ĐỒ ÁN
    p_sec3 = doc.add_paragraph()
    p_sec3.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_sec3.paragraph_format.space_before = Pt(6)
    p_sec3.paragraph_format.space_after = Pt(4)
    p_sec3.paragraph_format.line_spacing = 1.08
    r_s3 = p_sec3.add_run("3. KẾT QUẢ MONG MUỐN CUỐI ĐỒ ÁN")
    r_s3.font.name = 'Times New Roman'
    r_s3.font.size = Pt(12.5)
    r_s3.bold = True
    r_s3.font.color.rgb = RGBColor(0, 0, 0)

    bullets = [
        "Hoàn thành trọn vẹn Hệ thống Quản lý Bãi đỗ xe Thông minh (Smart Parking System) với 3 khu vực chuyên biệt: Customer Portal, Staff Portal và Admin Portal.",
        "Hoàn thành luồng nghiệp vụ khách hàng: Tìm kiếm bãi xe → Xem lưới ô đỗ trực quan → Chọn ô đỗ → Thuật toán A* dẫn đường mượt mà → Khóa ô đỗ chống xung đột (SELECT FOR UPDATE) → Đặt chỗ → Nhận vé điện tử QR Code → Đăng ký thẻ tháng.",
        "Hoàn thành luồng nghiệp vụ nhân viên và cổng kiểm soát: Quản lý bãi theo địa bàn phân công (assigned_lot) → Quét mã QR mở cần Barrier → Tiếp nhận xe vãng lai → Check-out tính phí tự động và xuất hóa đơn.",
        "Tích hợp mô hình Trí tuệ nhân tạo (Machine Learning) Random Forest Regressor 200 cây quyết định, dự báo chính xác tỷ lệ lấp đầy bãi xe trong 24 giờ tiếp theo hỗ trợ điều hành.",
        "Có đầy đủ bộ tài liệu phân tích kỹ thuật: Báo cáo đặc tả yêu cầu (SRS), Sơ đồ Use Case, ERD CSDL 6 bảng, Activity Diagram, Sequence Diagram, Wireframe và Kiến trúc 3 tầng.",
        "Có bộ kiểm thử tự động gồm 11 Unit Tests (Pytest), kiểm thử đa luồng 50 workers chống Race Condition và kịch bản Playwright E2E kiểm chứng minh họa thực tế.",
        "Hoàn thiện toàn bộ mã nguồn Backend FastAPI, Frontend Vue 3 + Tailwind CSS + GSAP 3, cơ sở dữ liệu PostgreSQL, kịch bản khởi chạy 1-click (run_web.bat) và báo cáo đồ án chuẩn học thuật sẵn sàng bảo vệ."
    ]

    for b in bullets:
        p_b = doc.add_paragraph()
        p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_b.paragraph_format.space_before = Pt(0)
        p_b.paragraph_format.space_after = Pt(2.5)
        p_b.paragraph_format.line_spacing = 1.08
        p_b.paragraph_format.left_indent = Cm(0.6)
        
        r_bullet = p_b.add_run("• ")
        r_bullet.font.name = 'Times New Roman'
        r_bullet.font.size = Pt(10.5)
        r_bullet.bold = True
        r_bullet.font.color.rgb = RGBColor(0, 0, 0)

        r_txt = p_b.add_run(b)
        r_txt.font.name = 'Times New Roman'
        r_txt.font.size = Pt(10.5)
        r_txt.font.color.rgb = RGBColor(0, 0, 0)

    # Lưu tệp an toàn (xử lý khi file đang mở trong WPS/Word)
    output_filename = "Ke_hoach_thuc_hien_10_tuan_SmartParkingSystem.docx"
    doc_path = os.path.join(os.getcwd(), output_filename)
    saved_path = doc_path
    try:
        doc.save(doc_path)
        print(f"==> Đã tạo kế hoạch 10 tuần thành công tại: {doc_path}")
    except PermissionError:
        base, ext = os.path.splitext(output_filename)
        saved_path = os.path.join(os.getcwd(), f"{base}_Nhom4{ext}")
        doc.save(saved_path)
        print(f"==> File gốc đang mở trong WPS/Word. Đã lưu thành: {saved_path}")

    # Copy vào thư mục Downloads của người dùng
    user_downloads = os.path.join(os.path.expanduser("~"), "Downloads")
    if os.path.exists(user_downloads):
        target_path = os.path.join(user_downloads, output_filename)
        try:
            shutil.copy2(saved_path, target_path)
            print(f"==> Đã sao chép sang thư mục Downloads: {target_path}")
        except PermissionError:
            base, ext = os.path.splitext(output_filename)
            alt_dl = os.path.join(user_downloads, f"{base}_Nhom4{ext}")
            shutil.copy2(saved_path, alt_dl)
            print(f"==> Downloads đang mở file gốc. Đã sao chép vào: {alt_dl}")

if __name__ == "__main__":
    build_smart_parking_plan()

