import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_presentation_docx():
    doc = Document()
    
    # Page setup - 0.75 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    # Color palette
    PRIMARY_COLOR = RGBColor(26, 54, 93)     # Deep Navy #1A365D
    SECONDARY_COLOR = RGBColor(43, 108, 176) # Slate Blue #2B6CB0
    ACCENT_COLOR = RGBColor(197, 48, 48)     # Accent Red #C53030
    DARK_TEXT = RGBColor(45, 55, 72)         # Charcoal #2D3748
    MUTED_TEXT = RGBColor(113, 128, 150)     # Gray #718096
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("KỊCH BẢN THUYẾT TRÌNH TIẾNG ANH (SLIDE 13 - 14 - 15)")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_COLOR

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("Đề tài: AI Model for Warning Students About Poor Study Posture\nPhenikaa University - GVHD: Mrs. Nguyễn Khánh Trâm")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = SECONDARY_COLOR
    run_sub.font.italic = True

    # Info Box (Table)
    info_table = doc.add_table(rows=1, cols=1)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = info_table.rows[0].cells[0]
    cell.width = Inches(7.0)
    set_cell_background(cell, "EDF2F7")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    p_info = cell.paragraphs[0]
    p_info.paragraph_format.space_after = Pt(2)
    r1 = p_info.add_run(" THÔNG TIN LƯỢT NÓI & NGUYÊN TẮC THIẾT KẾ:\n")
    r1.font.name = "Arial"
    r1.font.bold = True
    r1.font.size = Pt(10.5)
    r1.font.color.rgb = PRIMARY_COLOR
    
    r2 = p_info.add_run("• Vị trí người nói: Người thứ 5 (Thành viên chốt hạ toàn bài - Closer)\n"
                        "• Phạm vi phụ trách: Slide 13 (Conclusion), Slide 14 (Recommendations & Future Work), Slide 15 (Thank you & Q&A)\n"
                        "• Tổng thời lượng: Đúng 3 phút (khoảng 260 từ tiếng Anh ~ 85-90 từ/phút, cực kỳ thư thái, không sợ nói nhanh vấp từ)\n"
                        "• Đặc điểm lời thoại: Câu ngắn, từ vựng siêu dễ đọc, không dùng từ phức tạp líu lưỡi, có phiên âm tiếng Việt ngay bên dưới.")
    r2.font.name = "Arial"
    r2.font.size = Pt(10)
    r2.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def add_slide_section(slide_num, slide_title, duration, objective, body_language, speech_en, speech_vn, phonetics_list):
        # Header banner
        h_table = doc.add_table(rows=1, cols=1)
        h_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        h_cell = h_table.rows[0].cells[0]
        h_cell.width = Inches(7.0)
        set_cell_background(h_cell, "2B6CB0")
        set_cell_margins(h_cell, top=100, bottom=100, left=150, right=150)
        
        hp = h_cell.paragraphs[0]
        hp.paragraph_format.space_after = Pt(0)
        hrun = hp.add_run(f"SLIDE {slide_num}: {slide_title.upper()}  |  ⏱ Thời lượng: {duration}")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(11)
        hrun.font.bold = True
        hrun.font.color.rgb = RGBColor(255, 255, 255)
        
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        
        # Meta box (Mục tiêu + Cử chỉ)
        meta_table = doc.add_table(rows=1, cols=1)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        m_cell = meta_table.rows[0].cells[0]
        m_cell.width = Inches(7.0)
        set_cell_background(m_cell, "F7FAFC")
        set_cell_margins(m_cell, top=100, bottom=100, left=140, right=140)
        
        mp = m_cell.paragraphs[0]
        mp.paragraph_format.space_after = Pt(2)
        m_run1 = mp.add_run("🎯 Mục tiêu: ")
        m_run1.font.bold = True
        m_run1.font.name = "Arial"
        m_run1.font.size = Pt(9.5)
        m_run1.font.color.rgb = PRIMARY_COLOR
        
        m_run2 = mp.add_run(f"{objective}\n")
        m_run2.font.name = "Arial"
        m_run2.font.size = Pt(9.5)
        m_run2.font.color.rgb = DARK_TEXT
        
        m_run3 = mp.add_run("🎭 Cử chỉ & Điệu bộ: ")
        m_run3.font.bold = True
        m_run3.font.name = "Arial"
        m_run3.font.size = Pt(9.5)
        m_run3.font.color.rgb = ACCENT_COLOR
        
        m_run4 = mp.add_run(f"{body_language}")
        m_run4.font.name = "Arial"
        m_run4.font.size = Pt(9.5)
        m_run4.font.color.rgb = DARK_TEXT

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # English Speech
        p_en_title = doc.add_paragraph()
        p_en_title.paragraph_format.space_after = Pt(2)
        p_en_title_run = p_en_title.add_run("🗣️ LỜI THOẠI TIẾNG ANH (Đọc to, rõ, ngắt nghỉ theo từng gạch đầu dòng):")
        p_en_title_run.font.name = "Arial"
        p_en_title_run.font.size = Pt(10.5)
        p_en_title_run.font.bold = True
        p_en_title_run.font.color.rgb = PRIMARY_COLOR

        # English text box
        en_table = doc.add_table(rows=1, cols=1)
        en_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        e_cell = en_table.rows[0].cells[0]
        e_cell.width = Inches(7.0)
        set_cell_background(e_cell, "F0FFF4") # Light mint green
        set_cell_margins(e_cell, top=120, bottom=120, left=160, right=160)
        
        ep = e_cell.paragraphs[0]
        ep.paragraph_format.space_after = Pt(0)
        ep.paragraph_format.line_spacing = 1.25
        erun = ep.add_run(speech_en)
        erun.font.name = "Arial"
        erun.font.size = Pt(10.5)
        erun.font.color.rgb = DARK_TEXT

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # Vietnamese Phonics Table
        p_ph_title = doc.add_paragraph()
        p_ph_title.paragraph_format.space_after = Pt(2)
        p_ph_title_run = p_ph_title.add_run("📖 HƯỚNG DẪN PHÁT ÂM BỒI TIẾNG VIỆT (Tập đọc không lo vấp):")
        p_ph_title_run.font.name = "Arial"
        p_ph_title_run.font.size = Pt(10)
        p_ph_title_run.font.bold = True
        p_ph_title_run.font.color.rgb = SECONDARY_COLOR

        p_tab = doc.add_table(rows=len(phonetics_list) + 1, cols=3)
        p_tab.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        headers = ["Từ / Cụm từ tiếng Anh", "Phiên âm tiếng Việt dễ đọc", "Nghĩa tiếng Việt"]
        col_widths = [Inches(2.3), Inches(2.5), Inches(2.2)]
        for i, h in enumerate(headers):
            cell = p_tab.rows[0].cells[i]
            cell.width = col_widths[i]
            set_cell_background(cell, "E2E8F0")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            hp = cell.paragraphs[0]
            hrun = hp.add_run(h)
            hrun.font.name = "Arial"
            hrun.font.bold = True
            hrun.font.size = Pt(9)
            hrun.font.color.rgb = PRIMARY_COLOR
            
        for row_idx, (en_word, vi_phonetic, vi_meaning) in enumerate(phonetics_list, start=1):
            row = p_tab.rows[row_idx]
            for col_idx, text in enumerate([en_word, vi_phonetic, vi_meaning]):
                cell = row.cells[col_idx]
                cell.width = col_widths[col_idx]
                set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
                if row_idx % 2 == 0:
                    set_cell_background(cell, "F7FAFC")
                cp = cell.paragraphs[0]
                crun = cp.add_run(text)
                crun.font.name = "Arial"
                crun.font.size = Pt(9)
                if col_idx == 1:
                    crun.font.bold = True
                    crun.font.color.rgb = ACCENT_COLOR
                else:
                    crun.font.color.rgb = DARK_TEXT

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # Vietnamese meaning box
        p_vn_title = doc.add_paragraph()
        p_vn_title.paragraph_format.space_after = Pt(2)
        p_vn_title_run = p_vn_title.add_run("🇻🇳 NGHĨA TIẾNG VIỆT (Để bạn hiểu sâu từng câu mình đang nói):")
        p_vn_title_run.font.name = "Arial"
        p_vn_title_run.font.size = Pt(10)
        p_vn_title_run.font.bold = True
        p_vn_title_run.font.color.rgb = PRIMARY_COLOR

        vn_table = doc.add_table(rows=1, cols=1)
        vn_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        v_cell = vn_table.rows[0].cells[0]
        v_cell.width = Inches(7.0)
        set_cell_background(v_cell, "FFF5F5") # Light warm pink
        set_cell_margins(v_cell, top=100, bottom=100, left=140, right=140)
        
        vp = v_cell.paragraphs[0]
        vp.paragraph_format.space_after = Pt(0)
        vrun = vp.add_run(speech_vn)
        vrun.font.name = "Arial"
        vrun.font.size = Pt(9.5)
        vrun.font.color.rgb = DARK_TEXT
        vrun.font.italic = True

        doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Slide 13 Data
    s13_en = (
        '"Thank you, [Tên bạn trước].\n'
        'Hello teacher and everyone, I will now present the Conclusion and Future Work of our project.\n\n'
        'Please look at Slide 13. Our research achieved five main points:\n'
        '• First, Machine Learning performs much better than fixed Rule-Based methods, increasing accuracy by more than twenty-one percent.\n'
        '• Second, using a Temporal Split helps stop data leakage between video frames.\n'
        '• Third, our system can accurately classify five sitting postures, especially slouching with one hundred percent recall.\n'
        '• Fourth, Temporal Smoothing effectively cuts down false warnings from small body movements.\n'
        '• And fifth, the whole system can run smoothly in real time on normal computers."'
    )
    s13_vn = (
        'Cảm ơn bạn [Tên bạn trước]. Kính chào cô giáo và tất cả các bạn, sau đây em xin trình bày phần Kết luận và Hướng phát triển của dự án.\n\n'
        'Xin mời mọi người nhìn vào Slide 13. Nghiên cứu của chúng em đạt được năm kết quả chính:\n'
        '1. Thứ nhất, Máy học hoạt động vượt trội hơn hẳn phương pháp dùng luật cố định, nâng cao độ chính xác thêm hơn 21%.\n'
        '2. Thứ hai, việc chia dữ liệu theo thời gian (Temporal Split) giúp ngăn chặn triệt để rò rỉ dữ liệu giữa các khung hình.\n'
        '3. Thứ ba, hệ thống nhận diện chính xác 5 tư thế ngồi, đặc biệt là tư thế gù lưng đạt độ thu hồi tuyệt đối 100%.\n'
        '4. Thứ tư, cơ chế làm mượt theo thời gian (Temporal Smoothing) giúp loại bỏ hiệu quả các cảnh báo sai do cử động nhỏ.\n'
        '5. Và thứ năm, toàn bộ hệ thống hoạt động mượt mà theo thời gian thực trên máy tính thông thường.'
    )
    s13_phonetics = [
        ("Conclusion & Future Work", "Cần-clú-dừn èn Phiu-chơ Guớc", "Kết luận và Hướng phát triển"),
        ("Machine Learning", "Mơ-sin Lơ-ninh", "Học máy / Trí tuệ nhân tạo"),
        ("Rule-Based", "Ru-bây-st", "Phương pháp dựa trên quy tắc cố định"),
        ("Twenty-one percent (21%)", "Toen-ti uăn pơ-sừn", "Hai mươi mốt phần trăm"),
        ("Temporal Split", "Tém-pơ-rồ S-lít", "Chia tập dữ liệu theo thời gian"),
        ("Data leakage", "Đây-tơ lí-kịch", "Rò rỉ dữ liệu giữa train và test"),
        ("Accurately classify", "Ác-kiu-rớt-li c-lát-si-phai", "Phân loại một cách chính xác"),
        ("Slouching", "S-lao-chinh", "Tư thế ngồi gù / khom lưng"),
        ("Recall", "Ri-col", "Độ thu hồi / độ nhạy phát hiện"),
        ("Temporal Smoothing", "Tém-pơ-rồ S-mút-dinh", "Làm mượt dữ liệu theo thời gian"),
        ("False warnings", "Pho-s oóc-ninh-s", "Cảnh báo sai / cảnh báo giả"),
        ("Real time", "Riu-thaim", "Thời gian thực (không bị giật lag)")
    ]

    add_slide_section(
        slide_num=13,
        slide_title="Conclusion (Tổng kết kết quả nghiên cứu)",
        duration="~1 phút 15 giây (0:00 - 1:15)",
        objective="Khẳng định 5 thành tựu cốt lõi của đề tài (AI vượt trội hơn luật cứng, không rò rỉ dữ liệu, bắt gù lưng 100%, khử báo động giả, chạy real-time).",
        body_language="Nhìn sang bạn vừa nói xong, gật đầu mỉm cười nhẹ. Quay thẳng người về phía cô giáo và khán phòng. Khi nói 'First', 'Second' thì dùng ngón tay đếm nhịp hoặc chỉ tay lên màn hình slide.",
        speech_en=s13_en,
        speech_vn=s13_vn,
        phonetics_list=s13_phonetics
    )

    # Slide 14 Data
    s14_en = (
        '"Moving to Slide 14, we propose several recommendations for future work:\n\n'
        '• First, for data and AI: We will collect more data from more students, test in different classrooms, and improve the detection of difficult postures like uneven shoulders.\n\n'
        '• Second, for application: We plan to develop a friendly desktop or mobile app, and test the model on low-cost IoT devices.\n\n'
        '• Third, for user experience: We will add real-time visual and sound alerts to remind students directly, while keeping high speed and high accuracy."'
    )
    s14_vn = (
        'Chuyển sang Slide 14, nhóm chúng em đề xuất các định hướng phát triển trong tương lai:\n\n'
        '• Thứ nhất, về Dữ liệu và AI: Nhóm sẽ thu thập thêm mẫu từ nhiều học sinh hơn, thử nghiệm trong các lớp học thực tế khác nhau, và nâng cao độ chính xác ở các tư thế phức tạp như vai lệch.\n\n'
        '• Thứ hai, về Ứng dụng thực tế: Nhóm dự định phát triển ứng dụng tiện lợi trên máy tính hoặc điện thoại, đồng thời thử nghiệm mô hình trên các thiết bị IoT chi phí thấp.\n\n'
        '• Thứ ba, về Trải nghiệm người dùng: Hệ thống sẽ tích hợp cảnh báo bằng hình ảnh và âm thanh trực tiếp để nhắc nhở học sinh ngay lập tức, trong khi vẫn duy trì tốc độ xử lý cao và độ chính xác ổn định.'
    )
    s14_phonetics = [
        ("Recommendations", "Re-côm-mừn-đây-sừnz", "Đề xuất / khuyến nghị"),
        ("Different classrooms", "Đíp-phơ-rừn c-lát-s-rum-s", "Nhiều môi trường lớp học khác nhau"),
        ("Difficult postures", "Đíp-phi-cồ pót-s-chơ-s", "Các tư thế khó nhận diện"),
        ("Uneven shoulders", "Ăn-í-vừn sâu-đơ-s", "Tư thế vai lệch / không đều"),
        ("Develop an app", "Đi-ve-lốp ừn ép", "Phát triển phần mềm ứng dụng"),
        ("IoT devices", "Ai-âu-ti đi-vai-sịt-s", "Các thiết bị phần cứng IoT"),
        ("Visual and sound alerts", "Vi-ru-ồ èn sao ơ-lợt-s", "Cảnh báo bằng hình ảnh và âm thanh"),
        ("High accuracy", "Hai ác-kiu-rơ-si", "Độ chính xác cao")
    ]

    add_slide_section(
        slide_num=14,
        slide_title="Recommendations & Future Work (Định hướng tương lai)",
        duration="~1 phút 10 giây (1:15 - 2:25)",
        objective="Gom 7 ý trên slide thành 3 nhóm mạch lạc: (1) Dữ liệu & AI, (2) Phần mềm & Phần cứng IoT, (3) Cảnh báo âm thanh & Trải nghiệm thực tế.",
        body_language="Bấm remote chuyển slide dứt khoát. Giọng nói gợi mở, hào hứng về tương lai của sản phẩm. Tay mở rộng tự nhiên sang hai bên khi nói về 'friendly desktop or mobile app'.",
        speech_en=s14_en,
        speech_vn=s14_vn,
        phonetics_list=s14_phonetics
    )

    # Slide 15 Data
    s15_en = (
        '"To conclude, our project brings a practical, smart solution to protect students\' health in digital learning.\n\n'
        'On behalf of our team, I would like to thank our teacher, Mrs. Nguyen Khanh Tram, and all of you for listening.\n\n'
        'We are now ready and welcome any questions from you. Thank you very much!"'
    )
    s15_vn = (
        'Để kết luận lại, dự án của chúng em mang lại một giải pháp thông minh và thực tiễn nhằm bảo vệ sức khỏe cho học sinh trong môi trường học tập số hóa.\n\n'
        'Thay mặt cho toàn thể nhóm, em xin gửi lời cảm ơn sâu sắc nhất tới cô giáo hướng dẫn - cô Nguyễn Khánh Trâm, cùng toàn thể quý thầy cô và các bạn đã chú ý lắng nghe.\n\n'
        'Nhóm chúng em rất sẵn sàng và mong nhận được các câu hỏi, góp ý từ mọi người. Em xin chân thành cảm ơn!'
    )
    s15_phonetics = [
        ("To conclude", "Tu cần-clút", "Tóm lại / Để kết luận"),
        ("Smart solution", "S-mát sơ-lu-sừn", "Giải pháp thông minh"),
        ("Digital learning", "Đi-di-tồ lơ-ninh", "Môi trường học tập số"),
        ("On behalf of our team", "On bi-háp óv ao-ơ tim", "Thay mặt cho toàn đội ngũ"),
        ("Mrs. Nguyen Khanh Tram", "Mít-xịt Nguyễn Khánh Trâm", "Cô Nguyễn Khánh Trâm (GVHD)"),
        ("Welcome any questions", "Goeo-cơm en-ni quét-s-chừnz", "Nồng nhiệt đón nhận câu hỏi"),
        ("Thank you very much", "Thanh-kiu ve-ri mớt-ch", "Xin chân thành cảm ơn")
    ]

    add_slide_section(
        slide_num=15,
        slide_title="Thank You For Listening! (Lời cảm ơn & Q&A)",
        duration="~35 giây (2:25 - 3:00)",
        objective="Chốt hạ thông điệp nhân văn của dự án, tri ân cô giáo hướng dẫn và phong thái tự tin mời câu hỏi phản biện.",
        body_language="Chuyển sang slide robot dễ thương. Đứng thẳng, hai tay chắp nhẹ lịch sự trước bụng hoặc mở rộng đón nhận. Mắt nhìn thẳng vào cô giáo khi nhắc tên cô. Cúi chào nhẹ ở câu kết.",
        speech_en=s15_en,
        speech_vn=s15_vn,
        phonetics_list=s15_phonetics
    )

    # Bonus: Q&A Preparation Box
    doc.add_page_break()
    p_bonus_title = doc.add_paragraph()
    p_bonus_title.paragraph_format.space_after = Pt(6)
    r_bt = p_bonus_title.add_run("🎁 BÍ KÍP ỨNG PHÓ PHẦN HỎI ĐÁP (Q&A CHEAT-SHEET)")
    r_bt.font.name = "Arial"
    r_bt.font.size = Pt(14)
    r_bt.font.bold = True
    r_bt.font.color.rgb = PRIMARY_COLOR

    p_bonus_intro = doc.add_paragraph()
    p_bonus_intro.paragraph_format.space_after = Pt(8)
    r_bi = p_bonus_intro.add_run("Vì bạn là người kết bài, thầy cô rất dễ hỏi xoáy vào Slide 13 và Slide 14. Dưới đây là 3 câu hỏi 'tủ' hay gặp nhất kèm câu trả lời tiếng Anh cực ngắn, dễ nhớ giúp bạn ghi trọn điểm:")
    r_bi.font.name = "Arial"
    r_bi.font.size = Pt(10)
    r_bi.font.italic = True
    r_bi.font.color.rgb = MUTED_TEXT

    qa_list = [
        ("Câu 1: Why is Temporal Split better than Random Split?",
         "Tại sao dùng Temporal Split lại tốt hơn Random Split?",
         '"Because in video, consecutive frames are very similar. A temporal split prevents the model from seeing identical test frames during training, which avoids data leakage."',
         '"Bởi vì trong video, các khung hình cạnh nhau giống hệt nhau. Temporal split ngăn mô hình nhìn thấy khung hình test lúc huấn luyện, tránh được việc rò rỉ dữ liệu."'),
        
        ("Câu 2: How does Temporal Smoothing reduce false alarms?",
         "Làm sao cơ chế Temporal Smoothing giảm được cảnh báo nhầm?",
         '"We set a small time buffer, like 5 to 7 seconds. If a student just bends down quickly to pick up a pen, no alarm will ring. The alarm only sounds if the poor posture lasts for several seconds."',
         '"Nhóm đặt bộ đệm thời gian từ 5 đến 7 giây. Nếu học sinh chỉ cúi nhặt bút trong chốc lát, hệ thống sẽ không kêu. Chỉ khi tư thế xấu kéo dài liên tục thì mới báo động."'),
        
        ("Câu 3: What is the main limitation of your system right now?",
         "Hạn chế lớn nhất của hệ thống các bạn hiện nay là gì?",
         '"The main limit is occlusion: classroom desks hide the students\' legs, and a 2D camera cannot capture full 3D body depth. That is why we focus on upper-body features first."',
         '"Hạn chế lớn nhất là góc khuất: bàn học che mất chân học sinh và camera 2D không nhìn được chiều sâu 3D. Vì vậy nhóm tập trung vào nửa thân trên trước."')
    ]

    for q_en, q_vn, a_en, a_vn in qa_list:
        q_tab = doc.add_table(rows=1, cols=1)
        q_tab.alignment = WD_TABLE_ALIGNMENT.CENTER
        q_cell = q_tab.rows[0].cells[0]
        q_cell.width = Inches(7.0)
        set_cell_background(q_cell, "EDF2F7")
        set_cell_margins(q_cell, top=100, bottom=100, left=140, right=140)
        
        qp = q_cell.paragraphs[0]
        qp.paragraph_format.space_after = Pt(3)
        qrun1 = qp.add_run(f"❓ {q_en}\n")
        qrun1.font.name = "Arial"
        qrun1.font.bold = True
        qrun1.font.size = Pt(10)
        qrun1.font.color.rgb = PRIMARY_COLOR
        
        qrun2 = qp.add_run(f"   ({q_vn})\n\n")
        qrun2.font.name = "Arial"
        qrun2.font.size = Pt(9)
        qrun2.font.italic = True
        qrun2.font.color.rgb = DARK_TEXT

        qrun3 = qp.add_run(f"👉 Trả lời: {a_en}\n")
        qrun3.font.name = "Arial"
        qrun3.font.bold = True
        qrun3.font.size = Pt(9.5)
        qrun3.font.color.rgb = ACCENT_COLOR

        qrun4 = qp.add_run(f"   Dịch: {a_vn}")
        qrun4.font.name = "Arial"
        qrun4.font.size = Pt(9)
        qrun4.font.color.rgb = DARK_TEXT
        
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Save document
    filename = "Kich_Ban_Thuyet_Trinh_Slide_13_14_15.docx"
    doc.save(filename)
    print(f"File saved successfully as: {filename}")

if __name__ == "__main__":
    create_presentation_docx()
