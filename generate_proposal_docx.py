import os
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCX_PATH = os.path.abspath(r"D:\smart-parking-system\De_xuat_du_an_Smart_Parking_System.docx")
DOWNLOAD_DOCX = os.path.abspath(r"C:\Users\LENOVO\Downloads\De_xuat_du_an_Smart_Parking_System.docx")

doc = Document()

# Set standard academic margins (Top: 2.0cm, Bottom: 2.0cm, Left: 2.5cm, Right: 2.0cm)
for section in doc.sections:
    section.top_margin = Inches(0.79)     # ~2.0 cm
    section.bottom_margin = Inches(0.79)  # ~2.0 cm
    section.left_margin = Inches(0.98)    # ~2.5 cm
    section.right_margin = Inches(0.79)   # ~2.0 cm

# Configure Default Normal Style
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(13)
font.color.rgb = RGBColor(0, 0, 0)
style.paragraph_format.line_spacing = 1.25
style.paragraph_format.space_after = Pt(4)
style.paragraph_format.space_before = Pt(0)

def add_header_line(text, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13, space_after=3, space_before=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.2
    run = p.add_run(text)
    run.bold = bold
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    return p

def add_para(text="", bold_prefix="", bold=False, space_after=4, space_before=0, left_indent=0, justify=True):
    p = doc.add_paragraph()
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    if left_indent > 0:
        p.paragraph_format.left_indent = Inches(left_indent)
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(13)
    
    if text:
        r_text = p.add_run(text)
        r_text.bold = bold
        r_text.font.name = 'Times New Roman'
        r_text.font.size = Pt(13)
    return p

def add_bullet(text, level=1, bold_prefix=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.2
    
    indent = 0.28 * level
    p.paragraph_format.left_indent = Inches(indent)
    
    bullet_symbol = "●  " if level == 1 else ("○  " if level == 2 else "–  ")
    r_sym = p.add_run(bullet_symbol)
    r_sym.font.name = 'Times New Roman'
    r_sym.font.size = Pt(11 if level > 1 else 12)
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(13)
        
    r_text = p.add_run(text)
    r_text.font.name = 'Times New Roman'
    r_text.font.size = Pt(13)
    return p

# ==================== 1. TIÊU ĐỀ ĐỀ XUẤT ĐỀ TÀI ====================
add_header_line("CSE702131-1-1-26(N01) – ĐỒ ÁN CƠ SỞ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13, space_after=3)
add_header_line("ĐỀ XUẤT DỰ ÁN", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16, space_after=14)

# ==================== 2. THÔNG TIN DỰ ÁN ====================
add_para(" Hệ thống quản lý và dẫn đường bãi đỗ xe thông minh tích hợp thuật toán A* và mô hình học máy dự báo tỷ lệ lấp đầy (Smart Parking System - SPS)", bold_prefix="TÊN DỰ ÁN:", space_after=4, justify=False)
add_para(" Phạm Văn Đạt (MSV: 23017150)", bold_prefix="SINH VIÊN:", space_after=4, justify=False)
add_para(" 01/09/2026 – 30/10/2026", bold_prefix="THỜI GIAN THỰC HIỆN:", space_after=12, justify=False)

# ==================== 3. PHẠM VI CÔNG VIỆC ====================
add_para("", bold_prefix="PHẠM VI CÔNG VIỆC:", space_after=6, justify=False)

# Task 1
add_para(
    " Tiến hành nghiên cứu tổng quan về mô hình bài toán quản lý bãi đỗ xe thông minh (Smart Parking System) trong hạ tầng đô thị số (Smart City). Khảo sát nguyên lý hoạt động của các giải thuật tìm đường tối ưu trên lưới tọa độ 2D, trọng tâm là thuật toán A* (A-Star Search) kết hợp hàm ước lượng Heuristic khoảng cách Manhattan nhằm tìm lộ trình ngắn nhất và tự động né vật cản tĩnh/động. Nghiên cứu phương pháp học máy (Machine Learning) cho bài toán dự báo chuỗi thời gian tỷ lệ lấp đầy bãi xe (Occupancy Forecasting) sử dụng mô hình hồi quy Random Forest. Khảo sát kiến trúc backend API hiện đại với FastAPI, cơ chế phân quyền JWT Bearer Token, cơ sở dữ liệu quan hệ PostgreSQL và giải pháp kiểm soát đồng thời (Row-Level Locking) ngăn chặn xung đột đặt chỗ trùng lặp (Anti Double-Booking).",
    bold_prefix="Task 1: Research – Nghiên cứu & khảo sát công nghệ\n",
    space_after=8
)

# Task 2
add_para(
    " Tiến hành lập trình hoàn chỉnh hệ sinh thái phần mềm bằng ngôn ngữ Python, framework FastAPI, hệ quản trị cơ sở dữ liệu PostgreSQL, thư viện Scikit-learn và giao diện Vue 3 (Composition API) kết hợp Tailwind CSS. Hệ thống tập trung xây dựng và hoàn thiện các module nghiệp vụ cốt lõi sau:",
    bold_prefix="Task 2: Building – Xây dựng kiến trúc hệ thống, giải thuật tìm đường & mô hình AI\n",
    space_after=5
)

# Module 1
add_bullet("", level=1, bold_prefix="Kiến trúc Cơ sở dữ liệu & Phân hệ Backend Core (Database & Core Services)")
add_bullet("Thiết kế mô hình dữ liệu quan hệ (ERD) hoàn chỉnh gồm 6 bảng thực thể liên kết: users, parking_lots, parking_slots, bookings, sessions, occupancy_history; quản lý phiên bản lược đồ CSDL tự động bằng Alembic Migration.", level=2)
add_bullet("Xây dựng phân hệ Xác thực & Phân quyền bảo mật: Băm mật khẩu an toàn với Bcrypt, cấp phát và giải mã JWT Bearer Token, kiểm soát chặt chẽ quyền hạn truy cập của 3 vai trò (Quản trị viên - Admin, Nhân viên trực bãi - Staff, Khách hàng - Customer).", level=2)
add_bullet("Phân quyền nhân viên trực theo bãi đỗ cụ thể (Assigned Lot): Cho phép Admin phân bổ nhân viên về từng cơ sở (Đô Nghĩa, Vincom, Aeon Mall...) để quản lý barrier và kiểm soát xe vào/ra đúng khu vực phụ trách.", level=2)
add_bullet("Xây dựng nghiệp vụ Đặt chỗ đỗ xe (Booking), tự động sinh vé điện tử tích hợp mã QR Code (Base64) chứa mã tra cứu, biển số xe và thời gian gửi.", level=2)
add_bullet("Cài đặt cơ chế chống đặt chỗ trùng lặp (Anti Double-Booking) bằng kỹ thuật khóa cấp dòng Database (SELECT FOR UPDATE) triệt tiêu hoàn toàn nguy cơ tranh chấp tài nguyên (Race Condition) khi nhiều khách cùng thao tác.", level=2)
add_bullet("Xây dựng kênh truyền thông thời gian thực hai chiều WebSocket (ConnectionManager) phát sóng tức thời trạng thái các ô đỗ xe và telemetry cảm biến siêu âm IoT tới màn hình giám sát.", level=2)

# Module 2
add_bullet("", level=1, bold_prefix="Thuật toán Dẫn đường ngắn nhất A* (Pathfinding Module)")
add_bullet("Mô hình hóa không gian bãi đỗ xe thành ma trận lưới tọa độ 2D (Grid Matrix), nhận diện linh hoạt vật cản tĩnh (tường, cột trụ bê tông, đường biên) và vật cản động (các ô đang có xe đỗ).", level=2)
add_bullet("Cài đặt giải thuật tìm kiếm A* (astar.py) sử dụng hàng đợi ưu tiên (Priority Queue/heapq) và hàm ước lượng Heuristic Manhattan tối ưu khoảng cách di chuyển.", level=2)
add_bullet("Xây dựng dịch vụ định tuyến và endpoint API GET /api/lots/{lot_id}/navigate trả về chuỗi tọa độ (waypoints) dẫn từ cổng vào (0, 0) đến đúng ô đỗ được chọn với thời gian phản hồi tức thì (< 5ms).", level=2)

# Module 3
add_bullet("", level=1, bold_prefix="Pipeline Dữ liệu & Huấn luyện Mô hình Machine Learning (AI Engine)")
add_bullet("Sinh dữ liệu lịch sử giả lập (Data Generator): Xây dựng module sinh 4,320 mẫu chuỗi thời gian có quy luật thực tế phản ánh chu kỳ giờ cao điểm sáng/chiều, ngày trong tuần, cuối tuần, ngày lễ và độ nhiễu ngẫu nhiên Gaussian.", level=2)
add_bullet("Huấn luyện mô hình hồi quy (Model Training): Huấn luyện Random Forest Regressor (200 cây quyết định, max depth 10) từ dữ liệu lịch sử; tối ưu sai số tuyệt đối trung bình đạt độ chính xác cao (MAE < 0.06 ~ 5.8%).", level=2)
add_bullet("Đóng gói và cung cấp dịch vụ suy luận (Inference Service): Nạp mô hình Singleton (occupancy_model.joblib) phục vụ các API dự báo tỷ lệ lấp đầy theo từng khung giờ và dự báo xu hướng toàn bộ 24 giờ trong ngày.", level=2)

# Module 4
add_bullet("", level=1, bold_prefix="Giao diện Web điều hành tương tác thời gian thực (Frontend UI/UX)")
add_bullet("Xây dựng ứng dụng đơn trang (SPA) bằng Vue 3 (Composition API) kết hợp Tailwind CSS phong cách Dark Cyberpunk hiện đại, trực quan và tối ưu trải nghiệm người dùng.", level=2)
add_bullet("Trực quan hóa sơ đồ lưới bãi đỗ xe 2D hiển thị sinh động 4 trạng thái ô đỗ: Trống (Xanh), Đã đặt (Vàng), Đang có xe đỗ (Đỏ), Vật cản (Xám).", level=2)
add_bullet("Lập trình hiệu ứng chuyển động GSAP 3: Diễn hoạt xe ảo lăn bánh bám sát tọa độ đường đi A*, tự động xoay đầu xe theo góc cua rẽ mượt mà, hiệu ứng Stagger 2D và Numeric Counter số liệu thống kê.", level=2)
add_bullet("Xây dựng biểu đồ dự báo AI tương tác 24 giờ kèm thanh trượt scrubber trực quan; tích hợp chế độ Standalone Mock Engine tự động hỗ trợ chạy demo offline trên mọi trình duyệt.", level=2)

# Task 3
add_para(
    " Tiến hành kiểm thử toàn diện phần mềm nhằm đánh giá tính chính xác của thuật toán, độ ổn định của giao dịch dữ liệu và hiệu năng dự báo của mô hình AI:",
    bold_prefix="\nTask 3: Testing – Kiểm thử tự động & Đánh giá mô hình / hệ thống\n",
    space_after=4
)
add_bullet("Kiểm thử thuật toán A* (test_astar.py): Thực hiện 6 unit test cases tự động bằng Pytest kiểm tra đường thẳng không vật cản, tự động né vật cản, kịch bản bãi xe bị chặn hoàn toàn, điểm xuất phát trùng đích, độ dài đường đi tối ưu Manhattan.", level=2)
add_bullet("Kiểm thử mô hình Machine Learning (test_ml_engine.py): Thực hiện 5 unit test cases kiểm tra tính phân phối chuẩn của dữ liệu, tính tái lập theo seed ngẫu nhiên, quy luật giờ đêm thấp hơn giờ cao điểm, mô hình hội tụ và dự báo trong miền [0, 1].", level=2)
add_bullet("Kiểm thử kiểm soát đồng thời (Concurrency Test - test_double_booking.py): Mô phỏng đa luồng (multi-threading) 2 khách hàng cùng bấm đặt 1 ô đỗ tại cùng 1 mili-giây, kiểm chứng khóa hàng Row-level lock ngăn chặn triệt để lỗi Race Condition (hệ thống xử lý chuẩn xác: 1 khách thành công, 1 khách nhận mã 409 Conflict).", level=2)
add_bullet("Kiểm thử giao diện đầu-cuối (E2E Test): Kiểm thử tự động qua Chrome DevTools Protocol (CDP) trên trình duyệt Microsoft Edge headless kiểm tra phiên đăng nhập JWT, render sơ đồ lưới ô đỗ và giải phóng bộ nhớ.", level=2)

# Task 4
add_para(
    " Tiến hành tích hợp giao diện Frontend trực tiếp vào FastAPI server (phục vụ static files qua app.mount) cho phép toàn bộ hệ sinh thái chạy thống nhất trên một cổng dịch vụ duy nhất (http://localhost:8000). Đóng gói kịch bản khởi chạy 1 chạm cho môi trường Windows (run.bat và run.ps1) tự động kích hoạt môi trường ảo Python và máy chủ Uvicorn. Cấu hình hệ thống Swagger UI tự động tại /docs phục vụ kiểm thử và kết nối API chuẩn OpenAPI 3.0.",
    bold_prefix="Task 4: Deployment & Integration – Tích hợp Full-stack & Đóng gói luồng thực thi\n",
    space_after=8
)

# Task 5
add_para(
    " Tổng hợp toàn bộ số liệu kỹ thuật, báo cáo và tài liệu phục vụ quá trình nghiệm thu đề tài gồm:",
    bold_prefix="Task 5: Documentation & Presentation – Báo cáo và trình bày\n",
    space_after=4
)
add_bullet("Bảng số liệu phân tích tập dữ liệu lịch sử và biểu đồ đường cong tỷ lệ lấp đầy 24 giờ trong ngày.", level=1)
add_bullet("Báo cáo đánh giá hiệu năng thuật toán tìm đường ngắn nhất A* và so sánh độ tối ưu lộ trình.", level=1)
add_bullet("Báo cáo đánh giá độ chính xác của mô hình Machine Learning Random Forest (chỉ số MAE, RMSE).", level=1)
add_bullet("Cẩm nang tài liệu kỹ thuật toàn diện và hướng dẫn vận hành cho Khách hàng và Quản trị viên.", level=1)
add_bullet("Chuẩn bị slide thuyết trình và kịch bản chạy demo trực tiếp toàn bộ luồng hệ thống: Đăng nhập phân quyền -> Sơ đồ lưới 2D -> Thuật toán A* dẫn đường -> Đặt chỗ vé QR Code -> Check-in/out barrier -> Dự báo AI 24h.", level=1)

# ==================== 4. KẾT QUẢ DỰ ÁN ====================
add_para("", bold_prefix="\nKẾT QUẢ DỰ ÁN:", space_after=4, justify=False)
add_para("Các kết quả đạt được của dự án gồm:", space_after=4, justify=False)

# Nhóm 1: Báo cáo dự án
add_para("", bold_prefix="Báo cáo dự án", space_after=2, justify=False)
add_bullet("Tài liệu đề xuất dự án và nghiên cứu tổng quan hệ thống quản lý bãi đỗ xe thông minh.", level=1)
add_bullet("Tài liệu phân tích kiến trúc hệ thống, thiết kế CSDL quan hệ (ERD) và ma trận phân công trách nhiệm.", level=1)
add_bullet("Tài liệu nghiên cứu, cài đặt giải thuật tìm đường A* và mô hình học máy Random Forest dự báo công suất.", level=1)
add_bullet("Báo cáo kết quả kiểm thử tự động (Unit Tests, Concurrency Test) và đánh giá độ chính xác thực nghiệm.", level=1)

# Nhóm 2: Source Code
add_para("", bold_prefix="Source Code", space_after=2, justify=False)
add_bullet("Mã nguồn Python hoàn chỉnh cho phân hệ Backend FastAPI, cấu hình bảo mật JWT, SQLAlchemy ORM và Alembic.", level=1)
add_bullet("Mã nguồn module tìm đường ngắn nhất A* (astar.py, grid.py) và dịch vụ định tuyến navigation độc lập.", level=1)
add_bullet("Mã nguồn bộ sinh dữ liệu lịch sử giả lập (data_generator.py), module huấn luyện và suy luận AI (train.py, predict.py).", level=1)
add_bullet("Mã nguồn giao diện Web Frontend SPA (Vue 3, Tailwind CSS, GSAP Animation, WebSocket Client và Standalone Mock Engine).", level=1)
add_bullet("Bộ mã nguồn kiểm thử tự động gồm 11 unit test cases Pytest và script kiểm thử đa luồng chống Race Condition.", level=1)

# Nhóm 3: Sản phẩm thực nghiệm
add_para("", bold_prefix="Sản phẩm thực nghiệm", space_after=2, justify=False)
add_bullet("Hệ thống Web App hoàn chỉnh chạy tại http://localhost:8000 với đầy đủ giao diện cho Khách hàng và Ban quản lý.", level=1)
add_bullet("Mô hình Machine Learning Random Forest đã huấn luyện và tối ưu đạt sai số MAE < 0.06 (occupancy_model.joblib).", level=1)
add_bullet("Bộ kịch bản khởi chạy 1 chạm tự động (run.bat, run.ps1) giúp triển khai hệ thống dễ dàng trên máy tính Windows.", level=1)

# Nhóm 4: Tài liệu hỗ trợ
add_para("", bold_prefix="Tài liệu hỗ trợ", space_after=2, justify=False)
add_bullet("Slide báo cáo và thuyết trình dự án phục vụ bảo vệ đồ án trước hội đồng.", level=1)
add_bullet("Cẩm nang hướng dẫn sử dụng chi tiết cho Khách hàng và Người quản trị (User & Administrator Manual).", level=1)
add_bullet("Tài liệu đặc tả API Swagger UI tương tác trực quan tại đường dẫn /docs.", level=1)

# ==================== 5. CHỮ KÝ XÁC NHẬN ====================
p_foot = doc.add_paragraph()
p_foot.paragraph_format.space_before = Pt(28)
p_foot.paragraph_format.space_after = Pt(4)
p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r_date = p_foot.add_run("Hà Nội, Tháng 10/2026\n\n")
r_date.font.name = 'Times New Roman'
r_date.font.size = Pt(13)
r_date.italic = True

r_sign = p_foot.add_run("Xác nhận của GVHD\n\n\n\n")
r_sign.font.name = 'Times New Roman'
r_sign.font.size = Pt(13)
r_sign.bold = True

doc.save(DOCX_PATH)
print(f"SUCCESS: Saved DOCX at {DOCX_PATH}")

try:
    shutil.copy2(DOCX_PATH, DOWNLOAD_DOCX)
    print(f"SUCCESS: Copied to Downloads: {DOWNLOAD_DOCX}")
except Exception as e:
    print(f"Notice copy: {e}")
