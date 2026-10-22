import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document(r'C:\Users\LENOVO\Downloads\Ke_hoach_thuc_hien_10_tuan_BloomGift.docx')

print("=== SECTIONS ===")
for i, s in enumerate(doc.sections):
    print(f"Section {i}: Top={s.top_margin.cm:.2f}cm, Bottom={s.bottom_margin.cm:.2f}cm, Left={s.left_margin.cm:.2f}cm, Right={s.right_margin.cm:.2f}cm")
    print(f"Page width={s.page_width.cm:.2f}cm, Page height={s.page_height.cm:.2f}cm")

print("\n=== PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        fonts = [r.font.name for r in p.runs]
        sizes = [r.font.size.pt if r.font.size else None for r in p.runs]
        bolds = [r.bold for r in p.runs]
        colors = [r.font.color.rgb if r.font.color else None for r in p.runs]
        print(f"P{i}: Align={p.alignment}, Style={p.style.name}")
        print(f"   Text: {p.text}")
        print(f"   Runs: fonts={fonts}, sizes={sizes}, bolds={bolds}, colors={colors}")

print("\n=== TABLES ===")
for t_idx, table in enumerate(doc.tables):
    print(f"\n--- Table {t_idx+1} ---")
    print(f"Style: {table.style.name}, Alignment: {table.alignment}")
    for r_idx, row in enumerate(table.rows):
        cells_info = []
        for c in row.cells:
            p_info = " / ".join([p.text.strip() for p in c.paragraphs if p.text.strip()])
            cells_info.append(p_info)
        print(f"Row {r_idx}: {cells_info}")

