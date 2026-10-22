import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

def inspect_docx(path):
    print(f"\n==================================================")
    print(f"FILE: {path}")
    print(f"==================================================")
    doc = docx.Document(path)
    for i, s in enumerate(doc.sections):
        print(f"Section {i}: Top={s.top_margin.cm:.2f}cm, Bottom={s.bottom_margin.cm:.2f}cm, Left={s.left_margin.cm:.2f}cm, Right={s.right_margin.cm:.2f}cm")

    print("\n--- PARAGRAPHS ---")
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            runs_info = [(r.text[:30], r.font.name, r.font.size.pt if r.font.size else None, r.bold, r.italic) for r in p.runs]
            print(f"P{i}: Align={p.alignment}, Style={p.style.name}, SpaceAfter={p.paragraph_format.space_after}, SpaceBefore={p.paragraph_format.space_before}")
            print(f"   Text: {p.text}")
            print(f"   Runs: {runs_info}")

    print("\n--- TABLES ---")
    for t_idx, table in enumerate(doc.tables):
        print(f"\nTable {t_idx+1}: {len(table.rows)} rows, {len(table.columns)} cols, Align={table.alignment}")
        for r_idx, row in enumerate(table.rows):
            cells_info = []
            for c_idx, cell in enumerate(row.cells):
                tcPr = cell._tc.get_or_add_tcPr()
                shd = tcPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd')
                fill = shd.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill') if shd is not None else None
                cells_info.append(f"[{fill}] {cell.text.replace(chr(10), ' ').strip()[:35]}")
            print(f"  Row {r_idx}: {cells_info}")

if __name__ == "__main__":
    inspect_docx(r"C:\Users\LENOVO\Downloads\BienBanHopNhom_BloomGift.docx")
    inspect_docx(r"C:\Users\LENOVO\Downloads\Bien_ban_hop_nhom_BloomGift_Tuan_6_7_30-09-2026.docx")

