import docx
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

file_path = r"C:\Users\LENOVO\Downloads\BAO_CAO_DACS.docx"
print(f"File exists: {os.path.exists(file_path)}")
print(f"File size: {os.path.getsize(file_path):,} bytes")

doc = docx.Document(file_path)
print(f"Total sections: {len(doc.sections)}")
for i, s in enumerate(doc.sections):
    print(f"Section {i}: Page {s.page_width.cm:.2f}x{s.page_height.cm:.2f} cm | Margin T={s.top_margin.cm:.2f}, B={s.bottom_margin.cm:.2f}, L={s.left_margin.cm:.2f}, R={s.right_margin.cm:.2f} cm")

print(f"\nTotal paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

print("\n--- SAMPLE PARAGRAPHS (FIRST 40) ---")
for i, p in enumerate(doc.paragraphs[:40]):
    if p.text.strip():
        runs_info = [(r.text[:20], r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs[:2]]
        print(f"P{i:03d} [{p.style.name} | Align={p.alignment}]: {p.text[:80]} | {runs_info}")

print("\n--- CHAPTER HEADINGS IN DOC ---")
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if any(txt.upper().startswith(k) for k in ["CHƯƠNG", "MỤC LỤC", "LỜI MỞ ĐẦU", "LỜI CẢM ƠN", "LỜI CAM ĐOAN", "KẾT LUẬN", "TÀI LIỆU THAM KHẢO"]):
        print(f"P{i:03d}: {txt}")

print("\n--- TABLES SUMMARY ---")
for t_idx, t in enumerate(doc.tables[:10]):
    print(f"Table {t_idx+1}: {len(t.rows)} rows x {len(t.columns)} cols | Header: {[c.text.replace(chr(10), ' ').strip()[:20] for c in t.rows[0].cells]}")

